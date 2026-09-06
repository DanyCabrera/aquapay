# -*- coding: utf-8 -*-
"""
ERS rediseñada y resumida de AquaPay.
Alineada al sistema real (no al ERS desactualizado).
"""
from __future__ import annotations

import os
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DIAG = r"C:\Users\Usuario\Desktop\aquapay\docs\ers_diagramas"
OUT_DIRS = [
    r"c:\Users\Usuario\Desktop\Universidad Mariano Galvez 2022 Dany\Semestre #9\PROYECTO DE GRADUACIÓN",
    r"C:\Users\Usuario\Desktop\aquapay\docs",
]
OUT_NAME = "ERS-AquaPay.docx"

ACCENT = RGBColor(0x11, 0x11, 0x11)
MUTED = RGBColor(0x55, 0x55, 0x55)
WARN = RGBColor(0x9A, 0x34, 0x12)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.2)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11.5)):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = ACCENT
    st.paragraph_format.space_before = Pt(12)
    st.paragraph_format.space_after = Pt(6)


def shade(paragraph, hexcolor):
    pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    pr.append(shd)


def h1(t):
    doc.add_page_break()
    return doc.add_heading(t, level=1)


def h1_first(t):
    return doc.add_heading(t, level=1)


def h2(t):
    return doc.add_heading(t, level=2)


def h3(t):
    return doc.add_heading(t, level=3)


def p(text, bold=False, italic=False, color=None, center=False):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    return par


def bullet(text):
    par = doc.add_paragraph(text, style="List Bullet")
    par.paragraph_format.space_after = Pt(2)
    return par


def numbered(text):
    par = doc.add_paragraph(text, style="List Number")
    par.paragraph_format.space_after = Pt(2)
    return par


def note(text):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.font.size = Pt(10)
    run.italic = True
    run.font.color.rgb = MUTED
    shade(par, "F4F4F5")
    return par


def fig(path, caption, width=15.5):
    if os.path.exists(path):
        doc.add_picture(path, width=Cm(width))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p(f"[Diagrama no encontrado: {path}]", italic=True, color=WARN, center=True)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = MUTED
    return cap


def table(headers, rows, widths=None, caption=None, fs=8.5):
    if caption:
        c = doc.add_paragraph()
        r = c.add_run(caption)
        r.bold = True
        r.font.size = Pt(9)
        c.paragraph_format.space_before = Pt(6)
        c.paragraph_format.space_after = Pt(2)
    t = doc.add_table(rows=1, cols=len(headers))
    try:
        t.style = "Table Grid"
    except KeyError:
        pass
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        par = cell.paragraphs[0]
        run = par.add_run(h)
        run.bold = True
        run.font.size = Pt(fs)
        shade(par, "E8E8EA")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            run = par.add_run(str(val))
            run.font.size = Pt(fs)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def add_page_numbers():
    for s in doc.sections:
        footer = s.footer
        par = footer.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = par.add_run()
        f1 = OxmlElement("w:fldChar")
        f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = "PAGE"
        f2 = OxmlElement("w:fldChar")
        f2.set(qn("w:fldCharType"), "end")
        run._r.append(f1)
        run._r.append(it)
        run._r.append(f2)
        run.font.size = Pt(9)


def add_toc():
    p("Actualice el índice en Word: clic derecho → Actualizar campos → Actualizar toda la tabla.", italic=True)
    par = doc.add_paragraph()
    run = par.add_run()
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-2" \\h \\z \\u'
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "(Índice)"
    f2.append(t)
    f3 = OxmlElement("w:fldChar")
    f3.set(qn("w:fldCharType"), "end")
    run._r.append(f1)
    run._r.append(it)
    run._r.append(f2)
    run._r.append(f3)


# ===================== PORTADA =====================
for _ in range(2):
    doc.add_paragraph()
p("UNIVERSIDAD MARIANO GÁLVEZ DE GUATEMALA", bold=True, center=True)
p("Facultad de Ingeniería en Sistemas de Información y Ciencias de la Computación", center=True)
p("[Carrera — completar]", italic=True, center=True, color=MUTED)
for _ in range(2):
    doc.add_paragraph()
p("ESPECIFICACIÓN DE REQUISITOS DE SOFTWARE (ERS)", bold=True, center=True)
p("AquaPay: sistema web para la administración del cobro del servicio de agua potable", center=True)
p("Aldea Sibaná, El Asintal, Retalhuleu", center=True)
for _ in range(3):
    doc.add_paragraph()
p("Presentado por:", bold=True, center=True)
p("[Nombre del estudiante]", center=True, color=MUTED)
p("[Número de carné]", center=True, color=MUTED)
doc.add_paragraph()
p("Asesor:", bold=True, center=True)
p("[Nombre del asesor]", center=True, color=MUTED)
for _ in range(3):
    doc.add_paragraph()
p("Guatemala, [mes y año]", center=True, color=MUTED)

# ===================== INDICE =====================
doc.add_page_break()
h1_first("ÍNDICE")
add_toc()

# ===================== 1 =====================
h1("1. INTRODUCCIÓN")

h2("1.1 Propósito")
p(
    "Este documento especifica los requisitos de software de AquaPay, plataforma web "
    "para registrar beneficiarios y emitir recibos del servicio de agua potable del "
    "Comité de Agua Potable de Aldea Sibaná. Sirve como base para diseño, "
    "implementación, pruebas y evaluación del proyecto de graduación."
)

h2("1.2 Alcance del producto")
p(
    "AquaPay automatiza el padrón (beneficiarios, viviendas y chorros), el cálculo "
    "de deuda anual por chorro, la emisión de recibos correlativos con PDF y la "
    "consulta de indicadores/reportes. Opera con dos roles: administrador y tesorero."
)
p("Incluye:", bold=True)
bullet("Autenticación con DPI y contraseña (Supabase Auth + JWT).")
bullet("CRUD del padrón (solo administrador).")
bullet("Cobro de tarifa anual (uno o varios años) y compra de chorro (tesorero).")
bullet("PDF del recibo almacenado en Supabase Storage.")
bullet("Dashboard y reportes exportables a PDF.")
p("No incluye:", bold=True)
bullet("Pagos con tarjeta o pasarela en línea.")
bullet("Aplicación móvil nativa ni portal del beneficiario.")
bullet("MFA/TOTP (columnas reservadas; no implementado).")
bullet("Anulación de recibos ni modo offline.")

h2("1.3 Definiciones")
table(
    ["Término", "Definición"],
    [
        ["Beneficiario", "Vecino del servicio (usuarios_comunidad). No inicia sesión."],
        ["Chorro", "Punto de agua asociado a una vivienda; base del cobro anual."],
        ["Tarifa anual", "Monto fijo por chorro activo, configurable por el administrador."],
        ["Recibo", "Comprobante numerado (correlativo anual) con PDF."],
        ["DPI", "Documento Personal de Identificación (13 dígitos mínimo)."],
        ["ERS", "Especificación de Requisitos de Software."],
    ],
    widths=[4.0, 12.0],
    caption="Tabla 1. Definiciones",
)

h2("1.4 Referencias")
bullet("IEEE Std 830-1998 — Recommended Practice for Software Requirements Specifications.")
bullet("Código fuente AquaPay (frontend Next.js / backend Express / schema.sql).")
bullet("Documentación de Supabase Auth, Database y Storage.")

# ===================== 2 =====================
h1("2. DESCRIPCIÓN GENERAL")

h2("2.1 Perspectiva del producto")
p(
    "Aplicación web cliente-servidor: el frontend no accede a la base de datos; "
    "consume una API REST protegida con JWT. Auth, PostgreSQL y Storage se delegan "
    "en Supabase. Despliegue: Vercel (cliente), Render (API), Supabase (datos)."
)
fig(os.path.join(DIAG, "07_contexto.png"), "Figura 1. Diagrama de contexto", 14.5)

h2("2.2 Funciones principales")
table(
    ["Módulo", "Función"],
    [
        ["Autenticación", "Registro, login, perfil y cierre de sesión."],
        ["Padrón", "Beneficiarios, viviendas y chorros."],
        ["Facturación", "Pendientes, cobro anual, compra de chorro, PDF."],
        ["Configuración", "Tarifa anual vigente."],
        ["Reportes / Dashboard", "Indicadores e informes exportables."],
    ],
    widths=[4.5, 11.5],
    caption="Tabla 2. Funciones del sistema",
)

h2("2.3 Usuarios y roles")
table(
    ["Rol", "Responsabilidad", "Módulos"],
    [
        ["Administrador", "Padrón, tarifa, historial y reportes. No emite recibos.", "Dashboard, Beneficiarios, Pagos, Reportes, Configuración"],
        ["Tesorero", "Recibe pagos y emite recibos. No modifica el padrón.", "Dashboard, Facturas"],
    ],
    widths=[3.2, 6.0, 6.8],
    caption="Tabla 3. Roles (alineados al código)",
)
note(
    "Corrección respecto al ERS anterior: el rol operativo es «tesorero» (no «secretario»). "
    "MFA no es obligatorio en la implementación actual."
)

h2("2.4 Restricciones")
bullet("Requiere conexión a Internet (dependencia de Supabase/Vercel/Render).")
bullet("DPI de 13–20 dígitos; contraseña 8–72 caracteres.")
bullet("El importe del cobro se calcula en el servidor; el cliente no lo impone.")
bullet("Unicidad: (chorro_id, año) en pagos anuales; (año, secuencia) en recibos.")

h2("2.5 Supuestos")
bullet("Existe un comité con al menos un administrador y un tesorero.")
bullet("El cobro es anual por chorro (no mensual).")
bullet("El recibo impreso se firma de forma manuscrita.")

h2("2.6 Problema que resuelve")
p(
    "El control manual de cobros genera errores de cálculo, pérdida de información "
    "y dificultad para localizar beneficiarios. AquaPay centraliza el padrón, "
    "calcula la deuda acumulada y deja evidencia digital (recibo + PDF)."
)

# ===================== 3 =====================
h1("3. REQUISITOS FUNCIONALES")
p(
    "Numeración continua. Prioridad: Alta = núcleo del negocio; Media = soporte; "
    "Baja = mejora de experiencia."
)

h2("3.1 Autenticación")
table(
    ["ID", "Requisito", "Actor", "Pri."],
    [
        ["RF-01", "Iniciar sesión con DPI y contraseña; devolver JWT.", "Admin / Tesorero", "Alta"],
        ["RF-02", "Registrar cuenta con nombre, DPI, rol y contraseña. Admin solo si no existe otro activo.", "Visitante", "Alta"],
        ["RF-03", "Consultar perfil autenticado.", "Admin / Tesorero", "Alta"],
        ["RF-04", "Cerrar sesión en el cliente (limpieza de token local).", "Admin / Tesorero", "Media"],
        ["RF-05", "Restringir funciones según rol (servidor + interfaz).", "Sistema", "Alta"],
    ],
    widths=[1.6, 9.2, 3.2, 1.6],
    caption="Tabla 4. RF — Autenticación",
)

h2("3.2 Padrón de beneficiarios")
table(
    ["ID", "Requisito", "Actor", "Pri."],
    [
        ["RF-06", "Registrar beneficiario (nombre, DPI, teléfono opcional).", "Administrador", "Alta"],
        ["RF-07", "Listar y buscar beneficiarios por nombre o DPI.", "Admin (escritura) / consulta autenticada", "Alta"],
        ["RF-08", "Consultar detalle con viviendas y chorros.", "Administrador", "Alta"],
        ["RF-09", "Actualizar datos del beneficiario.", "Administrador", "Alta"],
        ["RF-10", "Activar / desactivar beneficiario (baja lógica).", "Administrador", "Alta"],
        ["RF-11", "Registrar y actualizar viviendas (dirección, año inicio cobro).", "Administrador", "Alta"],
        ["RF-12", "Registrar y actualizar chorros (cantidad, precio compra, estado).", "Administrador", "Alta"],
    ],
    widths=[1.6, 9.2, 3.2, 1.6],
    caption="Tabla 5. RF — Padrón",
)

h2("3.3 Facturación y pagos")
table(
    ["ID", "Requisito", "Actor", "Pri."],
    [
        ["RF-13", "Al seleccionar beneficiario, autocompletar datos y vivienda.", "Tesorero", "Alta"],
        ["RF-14", "Calcular años pendientes y monto por año (tarifa × chorros activos).", "Tesorero", "Alta"],
        ["RF-15", "Cobrar tarifa anual de uno o varios años en un recibo.", "Tesorero", "Alta"],
        ["RF-16", "Cobrar compra de chorro (una vez por chorro).", "Tesorero", "Alta"],
        ["RF-17", "Asignar número de recibo correlativo (reinicia cada año).", "Sistema", "Alta"],
        ["RF-18", "Permitir elegir la fecha de pago.", "Tesorero", "Media"],
        ["RF-19", "Expresar el total en letras (quetzales).", "Sistema", "Media"],
        ["RF-20", "Generar PDF, guardarlo en Storage y entregar descarga.", "Sistema", "Alta"],
        ["RF-21", "Consultar recibos e historial por beneficiario/año.", "Admin / Tesorero", "Alta"],
        ["RF-22", "Registrar el tesorero emisor (creado_por).", "Sistema", "Alta"],
    ],
    widths=[1.6, 9.2, 3.2, 1.6],
    caption="Tabla 6. RF — Facturación",
)
note(
    "Corrección: no hay pagos mensuales. El tipo de cobro es tarifa_anual o compra_chorro."
)

h2("3.4 Configuración, dashboard y reportes")
table(
    ["ID", "Requisito", "Actor", "Pri."],
    [
        ["RF-23", "Consultar y actualizar tarifa anual.", "Administrador", "Alta"],
        ["RF-24", "Mostrar dashboard (usuarios, viviendas, chorros, ingresos del mes).", "Admin / Tesorero", "Media"],
        ["RF-25", "Reportes: usuarios, viviendas, pagos e ingresos.", "Administrador", "Media"],
        ["RF-26", "Exportar reportes a PDF.", "Administrador", "Media"],
        ["RF-27", "Tema claro/oscuro en la interfaz.", "Admin / Tesorero", "Baja"],
    ],
    widths=[1.6, 9.2, 3.2, 1.6],
    caption="Tabla 7. RF — Configuración y reportes",
)

h2("3.5 Casos de uso")
fig(os.path.join(DIAG, "01_casos_uso.png"), "Figura 2. Diagrama de casos de uso", 15.5)

# ===================== 4 =====================
h1("4. REQUISITOS NO FUNCIONALES")
table(
    ["ID", "Categoría", "Requisito"],
    [
        ["RNF-01", "Seguridad", "Toda ruta protegida exige JWT válido (jose + JWKS)."],
        ["RNF-02", "Seguridad", "Contraseñas gestionadas por Supabase Auth (no en perfiles)."],
        ["RNF-03", "Seguridad", "Autorización por rol en servidor (requireRoles)."],
        ["RNF-04", "Seguridad", "Validación de entrada con DTO (whitelist)."],
        ["RNF-05", "Seguridad", "Helmet, CORS restringido a FRONTEND_URL."],
        ["RNF-06", "Integridad", "UNIQUE (chorro_id, año) evita doble cobro anual."],
        ["RNF-07", "Integridad", "Cobros dentro de transacción de base de datos."],
        ["RNF-08", "Rendimiento", "Búsqueda por DPI indexada; listados con límite."],
        ["RNF-09", "Usabilidad", "Interfaz en español, responsiva, mensajes claros."],
        ["RNF-10", "Disponibilidad", "Endpoint GET /health; depende de servicios en la nube."],
        ["RNF-11", "Compatibilidad", "Navegadores modernos; diseño responsive."],
        ["RNF-12", "Mantenibilidad", "TypeScript estricto; CI con lint/build/tsc."],
        ["RNF-13", "Datos", "PostgreSQL en Supabase; FK e índices en schema.sql."],
    ],
    widths=[1.8, 3.0, 11.2],
    caption="Tabla 8. Requisitos no funcionales",
)

# ===================== 5 =====================
h1("5. MODELO DE DATOS")
p(
    "Motor: PostgreSQL (Supabase). Nueve tablas. El cobro mensual y la columna "
    "cuota_mensual fueron descartados (migración 004)."
)
fig(os.path.join(DIAG, "05_er.png"), "Figura 3. Diagrama entidad-relación", 15.8)
table(
    ["Entidad", "Atributos clave", "Notas"],
    [
        ["perfiles", "id, nombre, dpi, rol, activo", "Cuentas del sistema"],
        ["usuarios_comunidad", "id, nombre_completo, dpi, telefono, activo", "Beneficiarios"],
        ["viviendas", "id, usuario_id, direccion, anio_inicio_cobro, activa", "Inicio de deuda"],
        ["chorros", "id, vivienda_id, precio_compra, cantidad, activo", "Sin cuota_mensual"],
        ["recibos", "numero_recibo, anio_recibo, secuencia, tipo_cobro, total", "Correlativo anual"],
        ["pagos_anuales", "recibo_id, chorro_id, anio, monto", "UNIQUE chorro+año"],
        ["pagos_compra_chorro", "recibo_id, chorro_id, monto, fecha", "Compra única"],
        ["recibos_pdf", "recibo_id, ruta_storage, url_publica", "1:1 con recibo"],
        ["configuracion", "clave, valor", "tarifa_anual, lugar_pago"],
    ],
    widths=[4.0, 7.5, 4.5],
    caption="Tabla 9. Entidades principales",
)

# ===================== 6 =====================
h1("6. ARQUITECTURA Y DESPLIEGUE")

h2("6.1 Stack tecnológico")
table(
    ["Capa", "Tecnología", "Uso"],
    [
        ["Frontend", "Next.js 15 + React 19 + Tailwind + shadcn", "Interfaz web"],
        ["Backend", "Node.js + Express + TypeORM + class-validator", "API REST"],
        ["Base de datos", "PostgreSQL (Supabase)", "Persistencia"],
        ["Auth", "Supabase Auth (JWT)", "Credenciales"],
        ["Storage", "Supabase Storage", "PDF de recibos"],
        ["PDF", "pdf-lib (+ node-signpdf opcional)", "Comprobantes"],
        ["Hosting", "Vercel + Render + Supabase", "Producción"],
        ["CI", "GitHub Actions", "Lint / build / tsc"],
    ],
    widths=[3.0, 7.0, 6.0],
    caption="Tabla 10. Stack",
)

h2("6.2 Arquitectura lógica")
fig(os.path.join(DIAG, "02_arquitectura.png"), "Figura 4. Arquitectura del sistema", 15.5)

h2("6.3 Vista de despliegue")
fig(os.path.join(DIAG, "06_despliegue.png"), "Figura 5. Vista de despliegue", 15.5)
table(
    ["Componente", "Plataforma", "Notas"],
    [
        ["Frontend", "Vercel", "Root: frontend"],
        ["API", "Render", "Root: backend · npm start"],
        ["DB / Auth / Storage", "Supabase", "Bucket privado recibos"],
    ],
    widths=[4.0, 4.0, 8.0],
    caption="Tabla 11. Despliegue",
)
note(
    "Corrección: no existe Realtime Gateway ni Redis en el proyecto actual."
)

# ===================== 7 =====================
h1("7. DIAGRAMAS DE COMPORTAMIENTO")

h2("7.1 Secuencia — emisión de recibo")
fig(os.path.join(DIAG, "03_secuencia.png"), "Figura 6. Diagrama de secuencia (tarifa anual)", 15.5)

h2("7.2 Flujo — facturación")
fig(os.path.join(DIAG, "04_flujo.png"), "Figura 7. Diagrama de flujo de facturación", 12.5)

h2("7.3 Regla de cálculo de deuda")
p(
    "Sea Ai el año de inicio de cobro de la vivienda y C el conjunto de chorros "
    "activos. Un año y ∈ [Ai … año actual] está pendiente si al menos un chorro de C "
    "no tiene registro en pagos_anuales para ese año. "
    "Monto por año = tarifa_anual × Σ(cantidad de chorros activos)."
)

# ===================== 8 =====================
h1("8. INTERFACES Y VALIDACIONES")

h2("8.1 Interfaces externas")
table(
    ["Interfaz", "Detalle"],
    [
        ["Usuario", "Web responsive (español); temas claro/oscuro."],
        ["API", "REST JSON bajo /api/auth, /api/usuarios, /api/facturacion, /api/reportes."],
        ["Auth", "Supabase Auth; verificación JWT con JWKS."],
        ["Storage", "PDF en bucket recibos; URL firmada temporal."],
    ],
    widths=[3.5, 12.5],
    caption="Tabla 12. Interfaces",
)

h2("8.2 Validaciones mínimas")
bullet("DPI: 13–20 dígitos, único en perfiles y en beneficiarios.")
bullet("Contraseña: 8–72 caracteres.")
bullet("Años de cobro: 2000–2100; deben figurar como pendientes.")
bullet("Propiedades no declaradas en DTO → rechazo (400).")

# ===================== 9 =====================
h1("9. REQUISITOS FUTUROS (NO IMPLEMENTADOS)")
bullet("MFA/TOTP obligatorio.")
bullet("Invalidación real del token al cerrar sesión.")
bullet("Anulación de recibos con justificación.")
bullet("Portal de consulta para el beneficiario.")
bullet("Notificaciones de morosidad.")
bullet("UNIQUE(chorro_id) en pagos_compra_chorro (hoy solo en aplicación).")

# ===================== 10 =====================
h1("10. CONCLUSIONES")
p(
    "Esta ERS define, de forma resumida y alineada a la implementación real de "
    "AquaPay, el problema, el alcance, los requisitos funcionales y no funcionales, "
    "el modelo de datos y la arquitectura. Corrige inconsistencias del documento "
    "previo (pagos mensuales, rol «secretario», MFA obligatorio y entidades obsoletas) "
    "y completa los diagramas de contexto, casos de uso, arquitectura, despliegue, "
    "secuencia, flujo y entidad-relación."
)
p(
    "El documento orienta las siguientes etapas del proyecto de graduación y facilita "
    "la trazabilidad entre requisitos, diseño e implementación."
)

# Auditoría breve
h2("Cambios realizados sobre el ERS anterior")
table(
    ["Hallazgo", "Corrección"],
    [
        ["Título «Propuesta» y sin estilos Heading", "Estructura ERS con títulos jerárquicos"],
        ["Saltos en RF (faltaban 10–14, 26–27)", "Numeración continua RF-01 a RF-27"],
        ["Rol Secretario / Administrados", "Administrador y Tesorero"],
        ["Pagos mensuales / pagos_mensuales", "Solo tarifa anual y compra de chorro"],
        ["MFA obligatorio", "Declarado como futuro / no implementado"],
        ["cuota_mensual en chorros", "Eliminada del modelo"],
        ["Realtime + Redis", "Retirado (no existe)"],
        ["3 diagramas", "7 diagramas incluidos"],
        ["Texto extenso / repetitivo", "Versión resumida por tablas"],
    ],
    widths=[6.5, 9.5],
    caption="Tabla 13. Mejoras de estructuración",
)

add_page_numbers()

paths = []
for folder in OUT_DIRS:
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, OUT_NAME)
    doc.save(path)
    paths.append(path)
    print("SAVED", path)

print("DONE")
