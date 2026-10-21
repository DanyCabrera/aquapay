"""Informe de QA manual del panel tesorero. No altera la aplicación."""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pathlib import Path

NAVY = RGBColor(0x14, 0x34, 0x59)
OUT = Path(__file__).resolve().parent / "Informe_QA_Panel_Tesorero_AquaPay.docx"


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def paragraph(doc, text, size=12, bold=False, align="left", space_after=8, italic=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    set_run_font(r, "Times New Roman", size, bold, italic, color)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, "Calibri", 14 if level == 1 else 12, True, color=NAVY)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        set_run_font(r, "Calibri", 9, True, color=RGBColor(255, 255, 255))
        shade_cell(cell, "143459")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.rows[ri + 1].cells[ci]
            cell.text = ""
            r = cell.paragraphs[0].add_run(val)
            set_run_font(r, "Calibri", 9)
            if ri % 2 == 1:
                shade_cell(cell, "F4F7FB")
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    sec.page_width = Cm(21.59)
    sec.page_height = Cm(27.94)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("AquaPay  ·  Aldea Sibaná"), "Calibri", 10, color=NAVY)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("Informe de QA  ·  Panel tesorero    "), "Calibri", 10, color=NAVY)
    add_page_number(footer)

    paragraph(doc, "INFORME DE QA MANUAL", size=14, bold=True, align="center", space_after=2)
    paragraph(doc, "Panel tesorero de AquaPay", size=16, bold=True, align="center", space_after=6)
    paragraph(
        doc,
        "Fecha: 20 de octubre de 2026. Ambiente: producción (https://aquapay-silk.vercel.app contra https://aquapay-production.up.railway.app). No se modificó código ni se hizo commit.",
        size=11,
        italic=True,
        align="center",
        space_after=12,
        color=RGBColor(0x55, 0x55, 0x55),
    )

    heading(doc, "1. Resumen")
    paragraph(
        doc,
        "El cobro en ventanilla funciona. El tesorero entra, emite tarifa anual y compra de chorro, abre el PDF, consulta el historial y no entra a las rutas del administrador. No se reprodujo ningún fallo alto que impida cobrar. Se reprodujeron 1 fallo medio y 3 fallos bajos. El restablecer contraseña por DPI, el 401 y el cierre de sesión sí devuelven a Entrar.",
        align="justify",
    )
    table(
        doc,
        ["Severidad", "Cantidad", "¿Bloquea el cobro?"],
        [
            ["Alta", "0", "No"],
            ["Media", "1", "No"],
            ["Baja", "3", "No"],
        ],
    )

    heading(doc, "2. Condiciones de la prueba")
    paragraph(
        doc,
        "Las cuentas de e2e (DPI 9990000000003) no existen en producción. Se creó una cuenta de tesorero de QA, DPI 8880 00000 0123, nombre QA Tesorero AquaPay, y se restableció su clave durante el recorrido. No se tocó la tarifa ni se borraron beneficiarios. Recibos emitidos en la prueba: 002 (tarifa anual, María Fernanda López García, Q 120.00, 20 oct 2026) y 003 (compra de chorro, la misma persona, Q 1,500.00, 20 oct 2026).",
        align="justify",
    )

    heading(doc, "3. Lo que funciona")
    bullets = [
        "Login: el botón no envía el formulario si el DPI no tiene 13 dígitos o la clave tiene menos de 8 caracteres. El DPI se muestra 4-5-4. Credenciales inválidas: «DPI o contraseña incorrectos.»",
        "Restablecer: DPI inexistente muestra «No hay una cuenta activa con ese DPI.» Claves distintas dejan Restablecer deshabilitado. Un cambio válido permite entrar con la clave nueva.",
        "Menú del tesorero: solo Dashboard, Emitir e Historial, más el acceso Emitir recibo. No aparecen Beneficiarios, Consultar pago, Reportes ni Tarifa.",
        "Dashboard: totales del mes en quetzales, tarifa vigente Q 120.00, accesos a Emitir e Historial y últimos recibos.",
        "Emitir tarifa: beneficiario con deuda, año 2026, total Q 120.00, letras «CIENTO VEINTE QUETZALES EXACTOS», correlativo 2026-002. Al guardar se abre el visor del recibo 002.",
        "Cobro duplicado de tarifa: al volver a elegir a la misma persona, «No hay tarifa anual pendiente» y Emitir queda deshabilitado.",
        "Emitir compra de chorro: Q 1,500.00, letras «MIL QUINIENTOS QUETZALES EXACTOS», recibo 003. Al intentar de nuevo: «Sin compra pendiente» y Emitir deshabilitado.",
        "El PDF se pide con token. Sin sesión, GET /api/facturacion/recibos/{id}/pdf responde 401 «Token requerido». El archivo del recibo 002 es un PDF válido (%PDF-1.7).",
        "Historial: listado, totales Q, búsqueda por nombre (Ana Lucía → 1 recibo), por número (002), año 2025 vacío (0 recibos), Ver PDF del recibo 001 (el que antes fallaba en disco) ahora abre. Anterior/Siguiente deshabilitados con una sola página.",
        "Autorización: /usuarios, /pagos, /reportes y /configuracion redirigen a /dashboard.",
        "Sesión: token inválido borra la sesión y vuelve a /login, sin overlay de error. Sin sesión, /facturacion pide login.",
    ]
    for b in bullets:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(b)
        set_run_font(r, "Times New Roman", 12)

    heading(doc, "4. Tabla de defectos")
    table(
        doc,
        ["Id", "Sev.", "Pantalla", "Pasos", "Resultado actual", "Resultado esperado", "¿Bloquea cobro?"],
        [
            [
                "QA-T01",
                "Media",
                "Dashboard",
                "Entrar como tesorero cuyo nombre es «QA Tesorero AquaPay».",
                "El saludo dice «Hola, QA.» Toma solo la primera palabra del nombre.",
                "Debe saludar con el nombre de ventanilla, no con un recorte que no identifica a la persona.",
                "No",
            ],
            [
                "QA-T02",
                "Baja",
                "Emitir",
                "Abrir Emitir con el listado de beneficiarios cargado.",
                "Texto «23 beneficiarios activos encontrados .» con un espacio antes del punto.",
                "«23 beneficiarios activos encontrados.»",
                "No",
            ],
            [
                "QA-T03",
                "Baja",
                "Historial",
                "Filtrar hasta dejar un solo recibo.",
                "Muestra «1 recibos».",
                "«1 recibo» cuando el total es uno.",
                "No",
            ],
            [
                "QA-T04",
                "Baja",
                "Emitir",
                "Emitir una compra de chorro y volver al formulario vacío.",
                "El tipo de pago sigue en «Compra de chorro». El correlativo sí avanza a 2026-004.",
                "Tras emitir, el formulario debería volver a tarifa anual, que es el cobro habitual de ventanilla.",
                "No",
            ],
        ],
    )

    heading(doc, "5. Orden sugerido de corrección")
    paragraph(
        doc,
        "1. QA-T01: el tesorero ve el dashboard en cada turno; un saludo cortado resta claridad. 2. QA-T04: no impide cobrar, pero deja el formulario en un concepto que ya no tiene deuda. 3. QA-T02 y QA-T03: copia, se pueden agrupar en un mismo arreglo de textos.",
        align="justify",
    )

    heading(doc, "6. Pruebas que no se pudieron completar")
    pending = [
        "Paginación de 12 en 12: en producción solo hay 3 recibos, así que Anterior y Siguiente permanecen deshabilitados. No se recorrió una segunda página.",
        "Vista previa visual del PDF: el visor abre un iframe con un PDF válido, pero en este navegador embebido el recuadro se vio negro. Descargar no falló. Hay que confirmar en Chrome o Edge de escritorio que las páginas se leen.",
        "Diálogo nativo de Imprimir: el clic no lanzó un cuadro de impresión comprobable en esta sesión.",
        "Beneficiario inactivo o sin vivienda: el listado de Emitir oculta inactivos (correcto). No se encontró en el padrón visible un caso sin vivienda ni sin chorro para forzar el mensaje de administrador.",
        "Búsqueda de historial usando solo el DPI: se probaron nombre y número de recibo. El DPI sí se muestra 4-5-4 en la tabla.",
        "Estado vacío total del dashboard: el mes ya tenía el recibo 001 y, tras la prueba, 002 y 003. No se vio un mes en ceros.",
    ]
    for b in pending:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(b)
        set_run_font(r, "Times New Roman", 12)

    heading(doc, "7. Observaciones (no son defectos de código)")
    paragraph(
        doc,
        "En Emitir aparecen beneficiarios de prueba llamados «E2E Cypress ADMIN…». No es un fallo del panel tesorero: son datos del padrón. Conviene que el administrador los retire si no son vecinos reales. El restablecer contraseña solo con el DPI es el diseño actual del sistema; no se reporta como error.",
        align="justify",
    )

    heading(doc, "8. Conclusión")
    paragraph(
        doc,
        "El panel tesorero en producción sirve para cobrar. Los cuatro defectos reproducidos son de texto y de estado del formulario; ninguno detiene la emisión del recibo ni la consulta del PDF autenticado. Lo primero a corregir es el saludo del dashboard.",
        align="justify",
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
