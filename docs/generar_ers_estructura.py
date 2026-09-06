# -*- coding: utf-8 -*-
"""ERS AquaPay — estructura académica solicitada."""
from __future__ import annotations

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DIAG = r"C:\Users\Usuario\Desktop\aquapay\docs\ers_diagramas"
OUT_PATHS = [
    r"C:\Users\Usuario\Desktop\aquapay\docs\ERS-AquaPay-v2.docx",
    r"c:\Users\Usuario\Desktop\Universidad Mariano Galvez 2022 Dany\Semestre #9\PROYECTO DE GRADUACIÓN\ERS-AquaPay-v2.docx",
]

INK = RGBColor(0x11, 0x11, 0x11)
MUTED = RGBColor(0x55, 0x55, 0x55)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.5)
sec.bottom_margin = Cm(2.5)
sec.left_margin = Cm(3.0)
sec.right_margin = Cm(2.5)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

for name, size in (("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 12), ("Heading 4", 12)):
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = INK
    st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
    st.paragraph_format.space_after = Pt(6)


def shade(par, color="E8E8EA"):
    pr = par._p.get_or_add_pPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), color)
    pr.append(sh)


def p(text, bold=False, italic=False, center=False, size=12, color=None):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    if color:
        run.font.color.rgb = color
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    return par


def h1(t):
    doc.add_page_break()
    return doc.add_heading(t, level=1)


def h2(t):
    return doc.add_heading(t, level=2)


def h3(t):
    return doc.add_heading(t, level=3)


def bullet(text):
    par = doc.add_paragraph(text, style="List Bullet")
    par.paragraph_format.space_after = Pt(2)
    return par


def note(text):
    par = p(text, italic=True, size=10, color=MUTED)
    shade(par, "F5F5F5")
    return par


def fig(filename, caption, width_cm=15.0):
    path = os.path.join(DIAG, filename)
    if os.path.exists(path):
        doc.add_picture(path, width=Cm(width_cm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p(f"[Insertar diagrama: {filename}]", italic=True, center=True, color=MUTED)
    p(caption, italic=True, center=True, size=10, color=MUTED)


def table(headers, rows, widths=None, caption=None, fs=9):
    if caption:
        c = doc.add_paragraph()
        r = c.add_run(caption)
        r.bold = True
        r.font.size = Pt(10)
        r.font.name = "Times New Roman"
        c.paragraph_format.space_before = Pt(8)
        c.paragraph_format.space_after = Pt(3)
    t = doc.add_table(rows=1, cols=len(headers))
    try:
        t.style = "Table Grid"
    except KeyError:
        pass
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        par = cell.paragraphs[0]
        run = par.add_run(str(h))
        run.bold = True
        run.font.size = Pt(fs)
        run.font.name = "Times New Roman"
        shade(par, "E8E8EA")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            run = par.add_run(str(val))
            run.font.size = Pt(fs)
            run.font.name = "Times New Roman"
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def rf(rid, nombre, desc, actor, pre, flujo, resultado):
    table(
        ["Campo", "Descripción"],
        [
            ["ID", rid],
            ["Nombre", nombre],
            ["Descripción", desc],
            ["Actor", actor],
            ["Precondiciones", pre],
            ["Flujo principal", flujo],
            ["Resultado esperado", resultado],
        ],
        widths=[3.5, 12.5],
        fs=9,
    )


def add_toc():
    p(
        "En Microsoft Word: clic derecho sobre el índice → Actualizar campos → Actualizar toda la tabla.",
        italic=True,
        size=10,
        color=MUTED,
    )
    par = doc.add_paragraph()
    run = par.add_run()
    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = 'TOC \\o "1-3" \\h \\z \\u'
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "(Índice)"
    f2.append(t)
    f3 = OxmlElement("w:fldChar")
    f3.set(qn("w:fldCharType"), "end")
    run._r.append(f1)
    run._r.append(it)
    run._r.append(f2)
    run._r.append(f3)


def add_page_numbers():
    for s in doc.sections:
        footer = s.footer
        par = footer.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = par.add_run()
        f1 = OxmlElement("w:fldChar")
        f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = "PAGE"
        f2 = OxmlElement("w:fldChar")
        f2.set(qn("w:fldCharType"), "end")
        run._r.append(f1)
        run._r.append(it)
        run._r.append(f2)
        run.font.size = Pt(10)


# ===================== PORTADA =====================
for _ in range(2):
    doc.add_paragraph()
p("UNIVERSIDAD MARIANO GÁLVEZ DE GUATEMALA", bold=True, center=True, size=14)
p("Facultad de Ingeniería en Sistemas de Información y Ciencias de la Computación", center=True)
p("[Carrera — completar]", italic=True, center=True, color=MUTED)
for _ in range(2):
    doc.add_paragraph()
p("ESPECIFICACIÓN DE REQUISITOS DE SOFTWARE (ERS)", bold=True, center=True, size=14)
p(
    "AquaPay: sistema web para la administración del cobro del servicio de agua potable",
    center=True,
)
p("Aldea Sibaná, El Asintal, Retalhuleu", center=True)
for _ in range(3):
    doc.add_paragraph()
p("Presentado por:", bold=True, center=True)
p("[Nombre completo del estudiante]", center=True, color=MUTED)
p("[Número de carné]", center=True, color=MUTED)
doc.add_paragraph()
p("Asesor:", bold=True, center=True)
p("[Nombre del asesor]", center=True, color=MUTED)
for _ in range(3):
    doc.add_paragraph()
p("Guatemala, [mes y año]", center=True, color=MUTED)

# ===================== ÍNDICE =====================
doc.add_page_break()
doc.add_heading("ÍNDICE", level=1)
add_toc()

# ===================== 1 =====================
h1("1. INTRODUCCIÓN")

h2("1.1 Propósito")
p(
    "El presente documento constituye la Especificación de Requisitos de Software (ERS) "
    "del sistema AquaPay. Su propósito es definir de manera clara, precisa y verificable "
    "los requisitos funcionales y no funcionales del sistema web destinado a la "
    "administración del cobro del servicio de agua potable del Comité de Agua Potable "
    "de Aldea Sibaná, municipio de El Asintal, departamento de Retalhuleu."
)
p(
    "Este documento está dirigido al estudiante desarrollador, al asesor del proyecto "
    "de graduación, al tribunal evaluador y a los futuros mantenedores del sistema. "
    "Sirve como base para el diseño, la implementación, las pruebas y la evaluación "
    "académica del producto."
)

h2("1.2 Alcance")
p(
    "AquaPay es una aplicación web que digitaliza el padrón de beneficiarios, el "
    "cálculo automático de la deuda anual por chorro, la emisión de recibos "
    "correlativos y la generación de comprobantes en formato PDF. El sistema opera "
    "con dos roles de acceso: administrador y tesorero."
)
p("Dentro del alcance se incluye:", bold=True)
bullet("Autenticación mediante DPI y contraseña.")
bullet("Gestión del padrón: beneficiarios, viviendas y chorros.")
bullet("Cobro de tarifa anual (uno o varios años) y cobro de compra de chorro.")
bullet("Numeración automática de recibos con reinicio anual.")
bullet("Generación y almacenamiento del recibo en PDF.")
bullet("Dashboard de indicadores y reportes exportables a PDF.")
bullet("Configuración de la tarifa anual vigente.")
p("Fuera del alcance:", bold=True)
bullet("Pagos en línea o pasarela de tarjetas.")
bullet("Aplicación móvil nativa.")
bullet("Portal de consulta para el beneficiario.")
bullet("Autenticación multifactor (MFA/TOTP), prevista pero no implementada.")
bullet("Anulación de recibos ya emitidos.")

h2("1.3 Definiciones, acrónimos y abreviaturas")
table(
    ["Término", "Definición"],
    [
        ["AquaPay", "Nombre del sistema web objeto de esta especificación."],
        ["Beneficiario", "Vecino que recibe el servicio de agua. No posee credenciales de acceso."],
        ["Chorro", "Punto de agua asociado a una vivienda; base del cobro anual."],
        ["DPI", "Documento Personal de Identificación (mínimo 13 dígitos)."],
        ["ERS", "Especificación de Requisitos de Software."],
        ["JWT", "JSON Web Token; credencial de sesión emitida por Supabase Auth."],
        ["PDF", "Portable Document Format; formato del comprobante de pago."],
        ["RBAC", "Control de acceso basado en roles."],
        ["Tarifa anual", "Monto fijo por chorro activo, configurable por el administrador."],
        ["Recibo", "Comprobante numerado que documenta un cobro realizado."],
        ["API", "Interfaz de programación de aplicaciones."],
        ["CRUD", "Crear, leer, actualizar y eliminar."],
    ],
    widths=[4.0, 12.0],
    caption="Tabla 1. Definiciones y acrónimos",
)

h2("1.4 Referencias")
bullet("IEEE Std 830-1998. Recommended Practice for Software Requirements Specifications.")
bullet("Código fuente del proyecto AquaPay (frontend Next.js y backend Express).")
bullet("Script de esquema: backend/src/database/schema.sql.")
bullet("Documentación oficial de Next.js, Express, TypeORM y Supabase.")
bullet("README.md del repositorio AquaPay.")

h2("1.5 Visión general del documento")
p(
    "El capítulo 2 describe el producto en términos generales e incluye la "
    "metodología de desarrollo. El capítulo 3 detalla los requisitos funcionales, "
    "no funcionales y de interfaces. El capítulo 4 presenta el modelo de datos y el "
    "diagrama entidad-relación. El capítulo 5 describe la arquitectura y el "
    "despliegue. El capítulo 6 agrupa los diagramas del sistema (contexto, casos de "
    "uso, secuencia, flujo, arquitectura, actividades y estados). El capítulo 7 "
    "delimita el alcance y las limitaciones. El capítulo 8 presenta las conclusiones."
)

# ===================== 2 =====================
h1("2. DESCRIPCIÓN GENERAL")

h2("2.1 Perspectiva del producto")
p(
    "AquaPay es un sistema web cliente-servidor independiente. El cliente (Next.js) "
    "no accede directamente a la base de datos; consume una API REST protegida con JWT. "
    "La autenticación, la base de datos PostgreSQL y el almacenamiento de archivos se "
    "delegan en Supabase. El frontend se despliega en Vercel y el backend en Render."
)
fig("07_contexto.png", "Figura 1. Diagrama de contexto del sistema AquaPay")

h2("2.2 Funciones del producto")
table(
    ["Módulo", "Función principal"],
    [
        ["Autenticación", "Registro, inicio de sesión, consulta de perfil y cierre de sesión."],
        ["Padrón", "CRUD de beneficiarios, viviendas y chorros."],
        ["Facturación", "Cálculo de deuda, cobro anual, compra de chorro y PDF."],
        ["Pagos", "Consulta del historial de recibos emitidos."],
        ["Reportes", "Informes de usuarios, viviendas, pagos e ingresos en PDF."],
        ["Configuración", "Actualización de la tarifa anual."],
        ["Dashboard", "Indicadores agregados de recaudación."],
    ],
    widths=[4.0, 12.0],
    caption="Tabla 2. Funciones del producto",
)

h2("2.3 Características de los usuarios")
table(
    ["Actor", "Descripción", "Módulos accesibles"],
    [
        [
            "Administrador",
            "Gestiona el padrón, la tarifa y los reportes. No emite recibos.",
            "Dashboard, Beneficiarios, Pagos, Reportes, Configuración",
        ],
        [
            "Tesorero",
            "Recibe pagos y emite recibos. No modifica el padrón.",
            "Dashboard, Facturas",
        ],
        [
            "Beneficiario",
            "Receptor del servicio. No es usuario del sistema.",
            "Ninguno (solo sujeto de registro y cobro)",
        ],
    ],
    widths=[3.2, 6.5, 6.3],
    caption="Tabla 3. Usuarios y roles del sistema",
)

h2("2.4 Restricciones")
bullet("El sistema requiere conexión permanente a Internet.")
bullet("Dependencia de Supabase para autenticación, base de datos y almacenamiento.")
bullet("DPI entre 13 y 20 dígitos; contraseña entre 8 y 72 caracteres.")
bullet("El monto del cobro se calcula exclusivamente en el servidor.")
bullet("Unicidad de pago anual por chorro: UNIQUE (chorro_id, anio).")
bullet("Correlativo de recibos único por año: UNIQUE (anio_recibo, secuencia).")
bullet("El primer administrador solo puede registrarse si no existe otro activo.")

h2("2.5 Supuestos y dependencias")
p("Supuestos:", bold=True)
bullet("Existe un comité con al menos un administrador y un tesorero capacitados.")
bullet("El esquema de cobro de la comunidad es anual por chorro, no mensual.")
bullet("El recibo impreso se firma de forma manuscrita por el tesorero.")
bullet("Los beneficiarios poseen un DPI válido y único.")
p("Dependencias:", bold=True)
bullet("Supabase (Auth, PostgreSQL, Storage).")
bullet("Vercel (hosting del frontend).")
bullet("Render (hosting del backend).")
bullet("Navegador web moderno (Chrome, Edge o Firefox).")

h2("2.6 Metodología a utilizar")
p(
    "Para el desarrollo de AquaPay se adopta una metodología incremental e "
    "iterativa, compatible con las prácticas ágiles, adaptada al alcance de un "
    "proyecto de graduación. Esta decisión se sustenta en la evidencia del "
    "repositorio: evolución del esquema mediante migraciones SQL numeradas, "
    "separación en capas (rutas, DTO, servicios, entidades) e integración continua "
    "con GitHub Actions."
)
p(
    "Las fases del ciclo de vida aplicadas al proyecto son las siguientes:"
)
table(
    ["Fase", "Actividades principales", "Producto / evidencia"],
    [
        [
            "1. Planificación",
            "Definición del problema, roles, alcance y reglas de cobro anual.",
            "README.md; reglas de negocio documentadas",
        ],
        [
            "2. Análisis",
            "Identificación de actores, requisitos funcionales/no funcionales y entidades del dominio.",
            "ERS; casos de uso; modelo conceptual",
        ],
        [
            "3. Diseño",
            "Arquitectura cliente-servidor, modelo ER, interfaces y contratos de API.",
            "Diagramas UML; schema.sql; stack tecnológico",
        ],
        [
            "4. Implementación",
            "Construcción incremental del frontend, backend y persistencia.",
            "Código en frontend/ y backend/",
        ],
        [
            "5. Pruebas",
            "Pruebas unitarias, verificación de tipos, lint y compilación automática.",
            "Vitest; GitHub Actions (CI)",
        ],
        [
            "6. Despliegue",
            "Publicación en la nube y configuración de variables de entorno.",
            "Vercel + Render + Supabase",
        ],
    ],
    widths=[3.2, 6.8, 6.0],
    caption="Tabla 3b. Fases de la metodología de desarrollo",
)
p(
    "Herramientas de apoyo a la metodología:",
    bold=True,
)
bullet("Control de versiones previsto mediante Git y flujo CI en GitHub Actions.")
bullet("Gestión del esquema con TypeORM (synchronize: false) y migraciones SQL explícitas.")
bullet("Validación continua de calidad: ESLint, tsc --noEmit y build de producción.")
note(
    "Advertencia: en el directorio de trabajo analizado no se localizó historial Git "
    "completo (.git). La metodología se documenta a partir de la estructura del "
    "proyecto y de los artefactos existentes; debe confirmarse con el cronograma "
    "oficial del anteproyecto."
)

# ===================== 3 =====================
h1("3. REQUISITOS ESPECÍFICOS")

h2("3.1 Requisitos funcionales")
p(
    "Los requisitos funcionales se derivaron de las pantallas, endpoints y reglas de "
    "negocio implementadas en el código fuente. Cada requisito es trazable a un "
    "módulo concreto del sistema."
)

h3("3.1.1 Módulo de autenticación")
rf(
    "RF-01",
    "Iniciar sesión",
    "Autenticar al usuario del sistema mediante DPI y contraseña, devolviendo un token JWT.",
    "Administrador, Tesorero",
    "El perfil debe existir y estar activo.",
    "1) El usuario ingresa DPI y contraseña. 2) El cliente valida formato. "
    "3) POST /api/auth/login. 4) El servidor verifica el perfil y delega la "
    "credencial en Supabase Auth. 5) Se retorna accessToken y datos del usuario.",
    "Código 200 con token y perfil. Ante error, 401 con mensaje genérico.",
)
rf(
    "RF-02",
    "Registrar cuenta de sistema",
    "Crear una cuenta con nombre, DPI, rol y contraseña. El rol administrador solo se concede si no existe otro activo.",
    "Visitante",
    "El DPI no debe estar registrado. Para administrador, conteo de activos = 0.",
    "1) GET /api/auth/registro-opciones. 2) El visitante completa el formulario. "
    "3) POST /api/auth/registro. 4) Se crea el usuario en Supabase Auth y el perfil.",
    "Código 201 con el perfil creado. 403 si el registro de administrador está cerrado.",
)
rf(
    "RF-03",
    "Consultar perfil autenticado",
    "Obtener los datos del usuario autenticado a partir del token.",
    "Administrador, Tesorero",
    "Token JWT válido.",
    "GET /api/auth/me con cabecera Authorization Bearer.",
    "Código 200 con id, nombre, dpi y rol.",
)
rf(
    "RF-04",
    "Cerrar sesión",
    "Eliminar la sesión local del navegador.",
    "Administrador, Tesorero",
    "Usuario autenticado.",
    "POST /api/auth/logout y limpieza de almacenamiento local.",
    "Sesión eliminada en el cliente. Nota: el token no se invalida en el servidor.",
)
rf(
    "RF-05",
    "Control de acceso por rol",
    "Restringir funcionalidades según el rol del perfil autenticado.",
    "Sistema",
    "Usuario autenticado con perfil activo.",
    "Middleware requireRoles en el servidor y control de rutas en el cliente.",
    "403 si el rol es insuficiente; redirección en la interfaz.",
)

h3("3.1.2 Módulo de gestión de usuarios / padrón")
rf(
    "RF-06",
    "Registrar beneficiario",
    "Crear un nuevo registro en el padrón de la comunidad.",
    "Administrador",
    "DPI único; nombre válido.",
    "POST /api/usuarios con nombre, DPI y teléfono opcional.",
    "Código 201 con el beneficiario creado.",
)
rf(
    "RF-07",
    "Listar y buscar beneficiarios",
    "Consultar el padrón con filtro opcional por nombre o DPI.",
    "Administrador",
    "Usuario autenticado.",
    "GET /api/usuarios?q=...",
    "Listado de beneficiarios coincidentes.",
)
rf(
    "RF-08",
    "Consultar detalle de beneficiario",
    "Mostrar beneficiario con sus viviendas y chorros.",
    "Administrador",
    "El identificador debe existir.",
    "GET /api/usuarios/:id",
    "Detalle completo del padrón del beneficiario.",
)
rf(
    "RF-09",
    "Actualizar beneficiario",
    "Modificar nombre, teléfono o estado del beneficiario.",
    "Administrador",
    "Beneficiario existente.",
    "PATCH /api/usuarios/:id",
    "Registro actualizado.",
)
rf(
    "RF-10",
    "Activar o desactivar beneficiario",
    "Realizar baja lógica sin eliminar el historial.",
    "Administrador",
    "Beneficiario existente.",
    "POST /api/usuarios/:id/activar o /desactivar",
    "Campo activo actualizado.",
)
rf(
    "RF-11",
    "Registrar y actualizar viviendas",
    "Asociar viviendas a un beneficiario con dirección y año de inicio de cobro.",
    "Administrador",
    "Beneficiario existente.",
    "POST /api/usuarios/:id/viviendas y PATCH /api/usuarios/viviendas/:id",
    "Vivienda creada o actualizada.",
)
rf(
    "RF-12",
    "Registrar y actualizar chorros",
    "Registrar puntos de agua con cantidad, precio de compra y estado.",
    "Administrador",
    "Vivienda existente.",
    "POST /api/usuarios/viviendas/:id/chorros y PATCH /api/usuarios/chorros/:id",
    "Chorro creado o actualizado.",
)

h3("3.1.3 Módulo de facturación")
rf(
    "RF-13",
    "Autocompletar datos al seleccionar beneficiario",
    "Al elegir un beneficiario, completar nombre, DPI, dirección y cantidad de chorros.",
    "Tesorero",
    "Beneficiario con al menos una vivienda.",
    "Selección en el formulario de facturación; consulta de detalle y pendientes.",
    "Formulario prellenado listo para cobro.",
)
rf(
    "RF-14",
    "Calcular años pendientes",
    "Determinar automáticamente los años adeudados y el monto por año.",
    "Tesorero",
    "Vivienda con chorros activos.",
    "GET /api/facturacion/pendientes/:viviendaId",
    "Lista de años pendientes, monto por año y chorros con compra pendiente.",
)
rf(
    "RF-15",
    "Cobrar tarifa anual",
    "Emitir un recibo que cancele uno o varios años de tarifa anual.",
    "Tesorero",
    "Los años enviados deben estar pendientes.",
    "POST /api/facturacion/cobrar/tarifa-anual dentro de una transacción.",
    "Código 201 con recibo y PDF. 400 si algún año no está pendiente.",
)
rf(
    "RF-16",
    "Cobrar compra de chorro",
    "Emitir recibo por el pago único de adquisición de un chorro.",
    "Tesorero",
    "Chorro sin pago de compra previo y con precio mayor a cero.",
    "POST /api/facturacion/cobrar/compra-chorro",
    "Código 201 con recibo y PDF. 409 si ya fue cobrada.",
)
rf(
    "RF-17",
    "Numeración automática de recibos",
    "Asignar correlativo con reinicio anual (001, 002, …).",
    "Sistema",
    "Operación de cobro en curso.",
    "Cálculo de siguiente secuencia del año actual.",
    "numero_recibo único; UNIQUE (anio_recibo, secuencia).",
)
rf(
    "RF-18",
    "Seleccionar fecha de pago",
    "Permitir al tesorero indicar la fecha del cobro.",
    "Tesorero",
    "Formulario de facturación abierto.",
    "Campo fecha en el formulario; se envía al servidor.",
    "fecha_pago registrada en el recibo.",
)
rf(
    "RF-19",
    "Monto en letras",
    "Convertir el total numérico a su expresión en quetzales.",
    "Sistema",
    "Total calculado.",
    "Utilidad de conversión a letras en servidor y vista previa en cliente.",
    "cantidad_en_letras almacenada e impresa en el PDF.",
)
rf(
    "RF-20",
    "Generar y almacenar PDF del recibo",
    "Crear el comprobante PDF, subirlo a Storage y entregarlo para descarga.",
    "Sistema",
    "Recibo persistido en transacción.",
    "pdf-lib + carga a bucket recibos + registro en recibos_pdf.",
    "PDF disponible en base64 y URL firmada temporal.",
)

h3("3.1.4 Módulo de pagos / historial")
rf(
    "RF-21",
    "Listar recibos",
    "Consultar recibos emitidos con búsqueda y límite.",
    "Administrador, Tesorero",
    "Usuario autenticado con rol permitido.",
    "GET /api/facturacion/recibos?q=&limit=",
    "Listado de recibos.",
)
rf(
    "RF-22",
    "Consultar historial por beneficiario",
    "Mostrar los recibos asociados a un beneficiario, con filtro opcional por año.",
    "Administrador",
    "Beneficiario existente.",
    "GET /api/facturacion/historial/:usuarioId",
    "Historial de pagos del beneficiario.",
)
rf(
    "RF-23",
    "Consultar detalle de un recibo",
    "Obtener un recibo específico con metadatos del PDF.",
    "Administrador, Tesorero",
    "Recibo existente.",
    "GET /api/facturacion/recibos/:id",
    "Detalle del recibo y URL del PDF.",
)

h3("3.1.5 Módulo de reportes")
rf(
    "RF-24",
    "Generar reportes administrativos",
    "Consultar reportes de usuarios, viviendas, pagos e ingresos.",
    "Administrador",
    "Usuario con rol administrador.",
    "GET /api/reportes/:tipo",
    "Tabla con columnas y filas del reporte solicitado.",
)
rf(
    "RF-25",
    "Exportar reportes a PDF",
    "Descargar el reporte seleccionado en formato PDF.",
    "Administrador",
    "Tipo de reporte válido.",
    "GET /api/reportes/:tipo/pdf",
    "Documento PDF descargable.",
)

h3("3.1.6 Módulo de configuración")
rf(
    "RF-26",
    "Consultar configuración",
    "Obtener la tarifa anual y el lugar de pago vigentes.",
    "Administrador, Tesorero",
    "Usuario autenticado.",
    "GET /api/facturacion/config",
    "Valores actuales de configuración.",
)
rf(
    "RF-27",
    "Actualizar tarifa anual",
    "Modificar el monto de la tarifa anual por chorro.",
    "Administrador",
    "Valor numérico válido.",
    "PUT /api/facturacion/config/tarifa",
    "Tarifa actualizada; se aplica a cobros posteriores.",
)
rf(
    "RF-28",
    "Consultar dashboard",
    "Mostrar indicadores de recaudación y totales del padrón.",
    "Administrador, Tesorero",
    "Usuario autenticado.",
    "GET /api/facturacion/dashboard",
    "Totales de usuarios, viviendas, chorros, recibos e ingresos.",
)
rf(
    "RF-29",
    "Tema claro u oscuro",
    "Permitir conmutar el tema visual de la interfaz.",
    "Administrador, Tesorero",
    "Interfaz cargada.",
    "Componente de tema con next-themes.",
    "Preferencia de tema aplicada y conservada en el navegador.",
)

h2("3.2 Requisitos no funcionales")

h3("3.2.1 Seguridad")
table(
    ["ID", "Requisito"],
    [
        ["RNF-01", "Toda ruta protegida debe exigir un JWT válido verificado contra el JWKS de Supabase."],
        ["RNF-02", "Las contraseñas no se almacenan en la base de datos de la aplicación; las gestiona Supabase Auth."],
        ["RNF-03", "El acceso a funciones sensibles debe restringirse por rol en el servidor."],
        ["RNF-04", "Las entradas deben validarse con DTO (whitelist)."],
        ["RNF-05", "El servidor debe aplicar Helmet y CORS restringido al origen del frontend."],
        ["RNF-06", "Las variables secretas deben residir en archivos .env excluidos del control de versiones."],
    ],
    widths=[2.0, 14.0],
    caption="Tabla 4. RNF de seguridad",
)

h3("3.2.2 Rendimiento")
table(
    ["ID", "Requisito"],
    [
        ["RNF-07", "Las búsquedas por DPI deben apoyarse en índices de base de datos."],
        ["RNF-08", "El listado de recibos debe limitar el tamaño de la respuesta (máximo 200 registros)."],
        ["RNF-09", "La generación del PDF del recibo debe completarse dentro de la misma operación de cobro."],
    ],
    widths=[2.0, 14.0],
    caption="Tabla 5. RNF de rendimiento",
)

h3("3.2.3 Usabilidad")
table(
    ["ID", "Requisito"],
    [
        ["RNF-10", "La interfaz debe estar completamente en idioma español."],
        ["RNF-11", "Los formularios deben validar datos antes del envío y mostrar errores por campo."],
        ["RNF-12", "La interfaz debe adaptarse a computadoras, tablets y teléfonos (diseño responsivo)."],
        ["RNF-13", "El sistema debe mostrar mensajes claros de éxito o error."],
    ],
    widths=[2.0, 14.0],
    caption="Tabla 6. RNF de usabilidad",
)

h3("3.2.4 Disponibilidad")
table(
    ["ID", "Requisito"],
    [
        ["RNF-14", "El backend debe exponer un endpoint de verificación de estado (GET /health)."],
        ["RNF-15", "La disponibilidad depende de los servicios en la nube (Vercel, Render, Supabase)."],
    ],
    widths=[2.0, 14.0],
    caption="Tabla 7. RNF de disponibilidad",
)

h3("3.2.5 Compatibilidad")
table(
    ["ID", "Requisito"],
    [
        ["RNF-16", "El sistema debe ejecutarse en Chrome, Edge y Firefox actualizados."],
        ["RNF-17", "No requiere instalación de software adicional en el equipo del usuario."],
    ],
    widths=[2.0, 14.0],
    caption="Tabla 8. RNF de compatibilidad",
)

h3("3.2.6 Mantenibilidad")
table(
    ["ID", "Requisito"],
    [
        ["RNF-18", "El código debe estar tipado con TypeScript en modo estricto."],
        ["RNF-19", "Debe existir separación en capas: rutas, DTO, servicios, entidades y middleware."],
        ["RNF-20", "La integración continua debe verificar lint, build y tipos en cada cambio."],
    ],
    widths=[2.0, 14.0],
    caption="Tabla 9. RNF de mantenibilidad",
)

h2("3.3 Requisitos de interfaces")

h3("3.3.1 Interfaces de usuario")
bullet("Interfaz web responsive construida con Next.js 15, React 19, Tailwind CSS y shadcn/ui.")
bullet("Navegación diferenciada: barra lateral para administrador; barra superior para tesorero.")
bullet("Formularios de login, registro, padrón, facturación, pagos, reportes y configuración.")
bullet("Soporte de tema claro y oscuro.")

h3("3.3.2 Interfaces de hardware")
p(
    "No se requieren dispositivos de hardware especializados. El sistema opera sobre "
    "computadoras, laptops, tablets o teléfonos con navegador web y acceso a Internet. "
    "La impresión del recibo PDF depende de una impresora convencional opcional."
)

h3("3.3.3 Interfaces de software")
table(
    ["Sistema externo", "Interfaz", "Uso"],
    [
        ["Supabase Auth", "API administrativa + JWKS", "Registro, login y verificación de JWT"],
        ["Supabase PostgreSQL", "Cadena DATABASE_URL (SSL)", "Persistencia de las 9 tablas"],
        ["Supabase Storage", "SDK Storage", "Almacenamiento de recibos PDF"],
        ["Cliente AquaPay", "HTTPS JSON /api/*", "Consumo de la API propia"],
    ],
    widths=[4.0, 5.0, 7.0],
    caption="Tabla 10. Interfaces de software",
)

h3("3.3.4 Interfaces de comunicaciones")
bullet("Protocolo: HTTPS en producción; HTTP en desarrollo local.")
bullet("Formato de intercambio: application/json.")
bullet("Autenticación: cabecera Authorization: Bearer <JWT>.")
bullet("Prefijos de API: /api/auth, /api/usuarios, /api/facturacion, /api/reportes.")

# ===================== 4 =====================
h1("4. MODELO DE DATOS")

h2("4.1 Descripción de entidades")
p(
    "El motor de base de datos es PostgreSQL, administrado mediante Supabase. "
    "El esquema se define en backend/src/database/schema.sql y comprende nueve tablas."
)
table(
    ["Entidad", "Propósito", "Atributos clave"],
    [
        ["perfiles", "Cuentas de acceso al sistema", "id, nombre, dpi, rol, activo"],
        ["usuarios_comunidad", "Beneficiarios del servicio", "id, nombre_completo, dpi, telefono, activo"],
        ["viviendas", "Domicilios del beneficiario", "id, usuario_id, direccion, anio_inicio_cobro"],
        ["chorros", "Puntos de agua", "id, vivienda_id, precio_compra, cantidad, activo"],
        ["recibos", "Comprobantes de cobro", "numero_recibo, anio_recibo, secuencia, tipo_cobro, total_pagado"],
        ["pagos_anuales", "Detalle de tarifa anual", "recibo_id, chorro_id, anio, monto_pagado"],
        ["pagos_compra_chorro", "Detalle de compra", "recibo_id, chorro_id, monto_pagado"],
        ["recibos_pdf", "Metadatos del PDF", "recibo_id, ruta_storage, url_publica"],
        ["configuracion", "Parámetros del sistema", "clave, valor (tarifa_anual, lugar_pago)"],
    ],
    widths=[3.8, 5.2, 7.0],
    caption="Tabla 11. Entidades del modelo de datos",
)
p("Restricciones de integridad relevantes:", bold=True)
bullet("UNIQUE (chorro_id, anio) en pagos_anuales — impide el doble cobro del mismo año.")
bullet("UNIQUE (anio_recibo, secuencia) y UNIQUE (numero_recibo) en recibos.")
bullet("tipo_cobro ∈ {tarifa_anual, compra_chorro}.")
bullet("Llaves foráneas con ON DELETE RESTRICT o CASCADE según la entidad.")

h2("4.2 Diagrama entidad-relación")
fig("05_er.png", "Figura 2. Diagrama entidad-relación de AquaPay", 15.5)
if os.path.exists(os.path.join(DIAG, "entity_relationship_diagram.png")):
    fig("entity_relationship_diagram.png", "Figura 2b. Diagrama ER exportado desde Mermaid", 15.5)

# ===================== 5 =====================
h1("5. ARQUITECTURA DEL SISTEMA")

h2("5.1 Arquitectura propuesta")
p(
    "AquaPay adopta una arquitectura de tres capas distribuidas: presentación "
    "(Next.js), lógica de negocio (Express + TypeORM) e infraestructura gestionada "
    "(Supabase). La comunicación entre cliente y servidor se realiza exclusivamente "
    "mediante HTTP/HTTPS sobre el prefijo /api."
)
fig("02_arquitectura.png", "Figura 3. Arquitectura lógica del sistema")

h2("5.2 Tecnologías utilizadas")
p(
    "A continuación se detallan las tecnologías empleadas en AquaPay, indicando "
    "la versión (según package.json), el propósito y el lugar concreto del proyecto "
    "donde se utiliza cada una."
)

h3("5.2.1 Resumen por capa")
table(
    ["Capa", "Tecnología", "Versión", "Uso en el proyecto"],
    [
        ["Frontend", "Next.js", "15.5.23", "Framework web; App Router en frontend/src/app"],
        ["Frontend", "React / React DOM", "19.1.0", "Componentes de interfaz en frontend/src/components"],
        ["Frontend", "TypeScript", "^5", "Tipado estático de todo el cliente"],
        ["Frontend", "Tailwind CSS", "^4", "Estilos utilitarios; frontend/src/app/globals.css"],
        ["Frontend", "shadcn/ui + @base-ui/react", "4.x / 1.7", "Componentes UI en frontend/src/components/ui"],
        ["Frontend", "next-themes", "^0.4.6", "Tema claro/oscuro (ThemeToggle)"],
        ["Frontend", "recharts", "^3.10.1", "Gráficos del dashboard / reportes"],
        ["Frontend", "sonner", "^2.0.8", "Notificaciones toast"],
        ["Frontend", "lucide-react", "^1.31.0", "Iconografía de la interfaz"],
        ["Frontend", "date-fns", "^4.4.0", "Formato y manipulación de fechas"],
        ["Backend", "Node.js", "22 (CI)", "Entorno de ejecución del servidor"],
        ["Backend", "Express", "^4.21.2", "API REST en backend/src (app.ts, routes/)"],
        ["Backend", "TypeScript", "^5.7.3", "Tipado del servidor"],
        ["Backend", "TypeORM", "^0.3.20", "ORM; entidades en backend/src/entities"],
        ["Backend", "pg", "^8.13.3", "Controlador PostgreSQL"],
        ["Backend", "class-validator / class-transformer", "0.14 / 0.5", "Validación de DTO en backend/src/dto"],
        ["Backend", "jose", "^5.9.6", "Verificación JWT/JWKS en middleware/auth"],
        ["Backend", "helmet / cors / morgan", "8 / 2.8 / 1.10", "Seguridad HTTP, CORS y logging"],
        ["Backend", "pdf-lib", "^1.17.1", "Generación de recibos PDF (pdf.service)"],
        ["Backend", "node-signpdf", "^3.0.0", "Firma digital P12 opcional del PDF"],
        ["Backend", "Vitest", "^3.0.5", "Pruebas unitarias (numero-a-letras)"],
        ["Datos / nube", "PostgreSQL (Supabase)", "gestionado", "Base de datos; schema.sql y migraciones"],
        ["Datos / nube", "Supabase Auth", "gestionado", "Login/registro; emisión de JWT"],
        ["Datos / nube", "Supabase Storage", "gestionado", "Bucket privado de recibos PDF"],
        ["Despliegue", "Vercel", "—", "Hosting del frontend"],
        ["Despliegue", "Render", "—", "Hosting del backend"],
        ["CI/CD", "GitHub Actions", "—", ".github/workflows/ci.yml"],
    ],
    widths=[2.4, 4.6, 2.4, 6.6],
    caption="Tabla 12. Tecnologías utilizadas por capa",
    fs=8,
)

h3("5.2.2 Dónde se utiliza cada tecnología")
table(
    ["Tecnología", "Dónde se utiliza", "Para qué"],
    [
        [
            "Next.js 15",
            "frontend/src/app/(auth), frontend/src/app/(app)",
            "Enrutamiento, páginas de login/registro, dashboard, padrón, facturación, reportes y configuración",
        ],
        [
            "React 19",
            "frontend/src/components/**",
            "Construcción de formularios, tablas, navegación y shells por rol",
        ],
        [
            "TypeScript",
            "frontend/ y backend/ (tsconfig strict)",
            "Tipado de contratos API, entidades, DTO y props de componentes",
        ],
        [
            "Tailwind CSS 4",
            "globals.css, className en páginas y componentes",
            "Diseño responsive, tema claro/oscuro y tipografía",
        ],
        [
            "shadcn/ui",
            "frontend/src/components/ui/",
            "Botones, inputs, diálogos, tablas, sheets y badges reutilizables",
        ],
        [
            "next-themes",
            "theme-provider / ThemeToggle",
            "Conmutar y persistir tema claro u oscuro",
        ],
        [
            "recharts",
            "Dashboard / reportes de ingresos",
            "Visualizar ingresos por mes",
        ],
        [
            "sonner",
            "layout raíz del frontend",
            "Mostrar mensajes de éxito y error al usuario",
        ],
        [
            "lucide-react",
            "navigation.ts y pantallas",
            "Iconos de menú y acciones",
        ],
        [
            "Express 4",
            "backend/src/app.ts, backend/src/routes/",
            "Exponer endpoints REST bajo /api/auth, /usuarios, /facturacion, /reportes",
        ],
        [
            "TypeORM",
            "backend/src/entities/, services/, data-source",
            "Mapear tablas, consultas y transacciones de cobro",
        ],
        [
            "class-validator",
            "backend/src/dto/ + middleware validate",
            "Validar DPI, contraseñas, años y cuerpos de petición",
        ],
        [
            "jose",
            "backend/src/middleware/auth.ts",
            "Verificar firma del JWT contra el JWKS de Supabase",
        ],
        [
            "helmet + cors",
            "backend/src/app.ts",
            "Cabeceras de seguridad y origen permitido (FRONTEND_URL)",
        ],
        [
            "pdf-lib",
            "backend/src/services/pdf.service.ts",
            "Generar el comprobante PDF del recibo",
        ],
        [
            "node-signpdf",
            "pdf.service (intentarFirmar)",
            "Firmar PDF si existe certificado P12 configurado",
        ],
        [
            "Supabase Auth",
            "auth.service.ts (createUser, signInWithPassword)",
            "Crear cuentas y autenticar con DPI/contraseña",
        ],
        [
            "PostgreSQL",
            "schema.sql, migraciones, TypeORM",
            "Persistir padrón, recibos, pagos y configuración",
        ],
        [
            "Supabase Storage",
            "facturacion.service (subir PDF)",
            "Guardar y firmar temporalmente las URL de los recibos",
        ],
        [
            "Vitest",
            "backend (npm test)",
            "Probar la conversión de montos a letras",
        ],
        [
            "Vercel",
            "Despliegue producción frontend",
            "Servir la interfaz Next.js",
        ],
        [
            "Render",
            "Despliegue producción backend",
            "Ejecutar la API Express",
        ],
        [
            "GitHub Actions",
            ".github/workflows/ci.yml",
            "Lint/build del frontend y verificación de tipos del backend",
        ],
    ],
    widths=[3.2, 5.8, 7.0],
    caption="Tabla 12b. Ubicación y propósito de cada tecnología",
    fs=8,
)

h3("5.2.3 Dependencias declaradas pero no utilizadas en código")
note(
    "En backend/package.json aparecen bcryptjs, otplib y qrcode; en frontend aparece "
    "@supabase/supabase-js. No se localizó uso efectivo de estas librerías en el "
    "código fuente analizado (la autenticación se delega en Supabase Auth y el "
    "cliente solo consume la API propia). Se documentan como dependencias residuales "
    "o previstas para MFA / cliente directo a Supabase."
)

h2("5.3 Vista de despliegue")
fig("06_despliegue.png", "Figura 4. Vista de despliegue")
table(
    ["Componente", "Plataforma", "Configuración"],
    [
        ["Frontend", "Vercel", "Root directory: frontend"],
        ["Backend / API", "Render", "Root: backend · build: npm run build · start: npm start"],
        ["Base de datos", "Supabase PostgreSQL", "Conexión SSL mediante DATABASE_URL"],
        ["Autenticación", "Supabase Auth", "Proveedor Email; JWT verificado por JWKS"],
        ["Archivos", "Supabase Storage", "Bucket privado recibos"],
    ],
    widths=[3.5, 4.0, 8.5],
    caption="Tabla 13. Componentes de despliegue",
)

# ===================== 6 =====================
h1("6. DIAGRAMAS")

h2("6.1 Diagrama de contexto")
p(
    "Representa al sistema AquaPay como una caja negra y sus interacciones con "
    "administradores, tesoreros y servicios externos de Supabase."
)
fig("07_contexto.png", "Figura 5. Diagrama de contexto")

h2("6.2 Diagrama de casos de uso")
p(
    "Muestra los casos de uso principales asociados a los actores administrador y "
    "tesorero, conforme a la navegación y a los endpoints implementados."
)
fig("01_casos_uso.png", "Figura 6. Diagrama de casos de uso")

h2("6.3 Diagrama de secuencia")
p(
    "Describe el flujo de emisión de un recibo de tarifa anual, desde la selección "
    "del beneficiario hasta la entrega del PDF."
)
fig("03_secuencia.png", "Figura 7. Diagrama de secuencia — emisión de recibo")

h2("6.4 Diagrama de flujo")
p(
    "Representa la lógica de decisión del proceso de facturación: verificación de "
    "deuda, elección del tipo de cobro, validación y emisión del recibo."
)
fig("04_flujo.png", "Figura 8. Diagrama de flujo — facturación", 12.5)

h2("6.5 Diagrama de arquitectura")
p(
    "Complementa el capítulo 5 mostrando las capas cliente, API, base de datos, "
    "almacenamiento y generación de PDF."
)
fig("02_arquitectura.png", "Figura 9. Diagrama de arquitectura (vista de capas)")

h2("6.6 Diagrama de actividades")
p(
    "El diagrama de actividades modela el flujo de trabajo del tesorero durante la "
    "emisión de un recibo, incluyendo decisiones de validación y la interacción con "
    "el sistema y con Supabase Storage."
)
fig("08_actividades.png", "Figura 10. Diagrama de actividades — emisión de recibo", 14.0)
p("Actividades principales representadas:", bold=True)
bullet("Seleccionar beneficiario y autocompletar datos / años pendientes.")
bullet("Decidir si existe deuda o compra pendiente.")
bullet("Elegir tipo de cobro (tarifa anual o compra de chorro).")
bullet("Capturar año/fecha y calcular el total automáticamente.")
bullet("Validar la operación; en caso de error, corregir datos.")
bullet("Emitir el recibo en transacción, almacenar el PDF y entregarlo al tesorero.")

h2("6.7 Diagrama de estados")
p(
    "El diagrama de estados describe los cambios de estado de los elementos clave "
    "del dominio: beneficiario, chorro, proceso de cobro (recibo) y sesión de usuario, "
    "según el comportamiento implementado en el sistema."
)
fig("09_estados.png", "Figura 11. Diagrama de estados de AquaPay", 15.5)
table(
    ["Objeto", "Estados", "Transiciones relevantes"],
    [
        [
            "Beneficiario",
            "Activo → Inactivo",
            "activar / desactivar (baja lógica; no se elimina el historial)",
        ],
        [
            "Chorro",
            "Activo / Inactivo; Compra pendiente → Compra pagada",
            "desactivar deja de generar deuda; cobrar compra cierra el pendiente",
        ],
        [
            "Recibo (cobro)",
            "Borrador → Validado → Emitido",
            "validar datos; emitir (estado terminal; no hay anulación)",
        ],
        [
            "Sesión",
            "No autenticado → Autenticado → Sesión cerrada",
            "login OK; logout o expiración del token",
        ],
    ],
    widths=[3.2, 5.5, 7.3],
    caption="Tabla 13b. Resumen de estados del sistema",
)
note(
    "El sistema no implementa anulación de recibos emitidos. Las columnas MFA en "
    "perfiles existen en el esquema, pero no alteran el ciclo de estados de sesión "
    "en la versión actual."
)

# Código Mermaid de apoyo (actividades y estados)
h3("6.8 Código Mermaid de los diagramas agregados")
p(
    "Los siguientes fragmentos permiten regenerar los diagramas de actividades y "
    "de estados en https://mermaid.live."
)
p("Diagrama de actividades:", bold=True)
code_m = doc.add_paragraph()
run = code_m.add_run(
    "flowchart TD\n"
    "  INICIO([Inicio]) --> SEL[Seleccionar beneficiario]\n"
    "  SEL --> AUTO[Autocompletar y cargar pendientes]\n"
    "  AUTO --> DEUDA{Hay deuda o compra?}\n"
    "  DEUDA -->|No| INFO[Informar sin pendiente] --> FIN1([Fin])\n"
    "  DEUDA -->|Si| TIPO[Elegir tipo de cobro]\n"
    "  TIPO --> DATOS[Año/fecha - total automatico]\n"
    "  DATOS --> VAL{Validacion OK?}\n"
    "  VAL -->|No| TIPO\n"
    "  VAL -->|Si| EMITIR[Emitir recibo + PDF]\n"
    "  EMITIR --> FIN2([Fin])"
)
run.font.name = "Consolas"
run.font.size = Pt(9)

p("Diagrama de estados (beneficiario):", bold=True)
code_m2 = doc.add_paragraph()
run2 = code_m2.add_run(
    "stateDiagram-v2\n"
    "  [*] --> Activo\n"
    "  Activo --> Inactivo: desactivar\n"
    "  Inactivo --> Activo: activar\n"
    "  Inactivo --> [*]: baja logica"
)
run2.font.name = "Consolas"
run2.font.size = Pt(9)

# ===================== 7 =====================
h1("7. ALCANCE Y LIMITACIONES")

h2("7.1 Funcionalidades incluidas")
bullet("Registro e inicio de sesión con DPI y contraseña.")
bullet("Control de acceso diferenciado para administrador y tesorero.")
bullet("Gestión completa del padrón (beneficiarios, viviendas y chorros).")
bullet("Cálculo automático de años pendientes y monto a pagar.")
bullet("Emisión de recibos por tarifa anual y por compra de chorro.")
bullet("Generación de PDF del recibo y almacenamiento en la nube.")
bullet("Dashboard de indicadores.")
bullet("Reportes administrativos exportables a PDF.")
bullet("Configuración de tarifa anual.")
bullet("Interfaz responsiva con tema claro/oscuro.")

h2("7.2 Fuera del alcance")
bullet("Pagos electrónicos o pasarelas de pago.")
bullet("Aplicación móvil nativa.")
bullet("App o portal para el beneficiario final.")
bullet("Notificaciones por correo o mensajería.")
bullet("Bitácora de auditoría de cambios.")
bullet("Anulación formal de recibos emitidos.")
bullet("Operación sin conexión a Internet.")

h2("7.3 Requisitos futuros")
bullet("Implementar MFA/TOTP obligatorio.")
bullet("Invalidar el token JWT al cerrar sesión.")
bullet("Agregar UNIQUE(chorro_id) en pagos_compra_chorro a nivel de base de datos.")
bullet("Anulación de recibos con justificación y trazabilidad.")
bullet("Consulta de estado de cuenta por DPI para el beneficiario.")
bullet("Notificaciones de morosidad.")
bullet("Activar la firma digital P12 del PDF en producción.")

# ===================== 8 =====================
h1("8. CONCLUSIONES")
p(
    "La presente Especificación de Requisitos de Software define de forma organizada "
    "las necesidades funcionales y no funcionales del sistema AquaPay, destinado a "
    "automatizar la administración del cobro del servicio de agua potable en Aldea "
    "Sibaná, El Asintal, Retalhuleu."
)
p(
    "A través del análisis del dominio y de la evidencia del código implementado, "
    "se establecieron veintinueve requisitos funcionales, veinte requisitos no "
    "funcionales, el modelo de datos de nueve entidades y la arquitectura "
    "cliente-servidor apoyada en Supabase, Vercel y Render."
)
p(
    "Este documento constituye la base para las etapas posteriores del proyecto de "
    "graduación: diseño detallado, implementación, pruebas y defensa ante el tribunal. "
    "Asimismo, facilita la trazabilidad entre lo requerido, lo diseñado y lo "
    "construido, reduciendo ambigüedades durante el desarrollo y el mantenimiento."
)
p(
    "Finalmente, la delimitación explícita del alcance y de los requisitos futuros "
    "permite distinguir con claridad lo implementado de las mejoras pendientes, "
    "contribuyendo a una evaluación académica transparente y coherente con el "
    "producto real."
)

# ===================== ANEXOS =====================
h1("ANEXOS")
h2("Anexo A. Resumen de requisitos funcionales")
table(
    ["ID", "Módulo", "Nombre"],
    [
        ["RF-01", "Autenticación", "Iniciar sesión"],
        ["RF-02", "Autenticación", "Registrar cuenta"],
        ["RF-03", "Autenticación", "Consultar perfil"],
        ["RF-04", "Autenticación", "Cerrar sesión"],
        ["RF-05", "Autenticación", "Control de acceso por rol"],
        ["RF-06", "Padrón", "Registrar beneficiario"],
        ["RF-07", "Padrón", "Listar y buscar"],
        ["RF-08", "Padrón", "Consultar detalle"],
        ["RF-09", "Padrón", "Actualizar beneficiario"],
        ["RF-10", "Padrón", "Activar / desactivar"],
        ["RF-11", "Padrón", "Viviendas"],
        ["RF-12", "Padrón", "Chorros"],
        ["RF-13", "Facturación", "Autocompletar datos"],
        ["RF-14", "Facturación", "Calcular pendientes"],
        ["RF-15", "Facturación", "Cobrar tarifa anual"],
        ["RF-16", "Facturación", "Cobrar compra de chorro"],
        ["RF-17", "Facturación", "Numeración de recibos"],
        ["RF-18", "Facturación", "Fecha de pago"],
        ["RF-19", "Facturación", "Monto en letras"],
        ["RF-20", "Facturación", "PDF del recibo"],
        ["RF-21", "Pagos", "Listar recibos"],
        ["RF-22", "Pagos", "Historial por beneficiario"],
        ["RF-23", "Pagos", "Detalle de recibo"],
        ["RF-24", "Reportes", "Generar reportes"],
        ["RF-25", "Reportes", "Exportar PDF"],
        ["RF-26", "Configuración", "Consultar configuración"],
        ["RF-27", "Configuración", "Actualizar tarifa"],
        ["RF-28", "Dashboard", "Indicadores"],
        ["RF-29", "Interfaz", "Tema claro/oscuro"],
    ],
    widths=[2.2, 3.5, 10.3],
    caption="Tabla 14. Catálogo resumido de RF",
)

h2("Anexo B. Código Mermaid del diagrama ER")
p(
    "El siguiente código puede pegarse en https://mermaid.live para regenerar el "
    "diagrama entidad-relación:"
)
code = doc.add_paragraph()
run = code.add_run(
    "erDiagram\n"
    "  PERFILES ||--o{ RECIBOS : emite\n"
    "  USUARIOS_COMUNIDAD ||--o{ VIVIENDAS : posee\n"
    "  VIVIENDAS ||--o{ CHORROS : tiene\n"
    "  VIVIENDAS ||--o{ RECIBOS : recibe\n"
    "  RECIBOS ||--o{ PAGOS_ANUALES : detalla\n"
    "  RECIBOS ||--o{ PAGOS_COMPRA_CHORRO : detalla\n"
    "  RECIBOS ||--|| RECIBOS_PDF : genera\n"
    "  CHORROS ||--o{ PAGOS_ANUALES : cobrado_en\n"
    "  CHORROS ||--o{ PAGOS_COMPRA_CHORRO : adquirido_en\n"
    "  PERFILES ||--o{ CONFIGURACION : actualiza"
)
run.font.name = "Consolas"
run.font.size = Pt(9)

add_page_numbers()

for path in OUT_PATHS:
    folder = os.path.dirname(path)
    if folder and not os.path.exists(folder):
        print("SKIP (folder missing):", path)
        continue
    doc.save(path)
    print("SAVED", path)

print("DONE")
