from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parent/'Manual_de_Usuario_Administrador_AQUAPAY.docx'
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
for s in doc.styles:
 for b in list(s.element.findall('.//'+qn('w:pBdr'))):b.getparent().remove(b)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(8)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
doc.styles['Title'].font.size=Pt(30);doc.styles['Heading 1'].font.size=Pt(20);doc.styles['Heading 2'].font.size=Pt(13);doc.styles['Caption'].font.size=Pt(9)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;foot.add_run('AQUAPAY  |  Administración  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
doc.core_properties.title='Manual de usuario del panel de Administrador de AQUAPAY';doc.core_properties.author='AQUAPAY'
def p(t):return doc.add_paragraph(t)
def img(name,caption,crop=False):
 f=ROOT/(name+'.png')
 if name=='acceso':f=ROOT.parent/'manual_assets'/'01_acceso_detalle.png'
 if crop:
  im=Image.open(f).crop((250,75,1440,960));f=ROOT/(name+'_detalle.png');im.save(f)
 w,h=Image.open(f).size;width=min(7.1,4.5*w/h)
 x=doc.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(4);x.add_run().add_picture(str(f),width=Inches(width));doc.inline_shapes[-1]._inline.docPr.set('descr',caption);doc.add_paragraph(caption,'Caption')
def page(title,lead,name=None,steps=(),note='',crop=False):
 doc.add_page_break();doc.add_heading(title,1);p(lead)
 if name:img(name,'Figura '+str(len(doc.inline_shapes)+1)+'. '+lead,crop)
 for i,t in enumerate(steps,1):
  x=p('');x.add_run(f'{i}. ').bold=True;x.add_run(t)
 if note:
  x=p('');x.add_run('Tenga en cuenta ').bold=True;x.add_run(note)

doc.add_paragraph('AQUAPAY','Subtitle');doc.add_paragraph('Manual de usuario del panel de Administrador','Title')
p('Guía para gestionar el padrón de beneficiarios, viviendas y chorros, consultar pagos, generar reportes y actualizar la tarifa anual del servicio de agua.')
p('Versión 1.0  |  27 de octubre de 2026')
doc.add_heading('Cómo utilizar este manual',1)
p('Siga la tarea que necesita realizar. Los círculos rojos de las capturas identifican los controles mencionados entre paréntesis. Los nombres, DPI e importes mostrados son ejemplos; en su sesión aparecerá la información de su comunidad.')
p('Necesita la dirección de AQUAPAY, conexión a internet y una cuenta activa con rol de administrador. El registro de cobros y la emisión de nuevos recibos corresponden al panel de Tesorería.')
doc.add_heading('Contenido',1)
for t in ['Acceso y resumen del panel · páginas 2 y 3','Búsqueda y exportación del padrón · página 4','Registro de persona vivienda y chorro · páginas 5 a 7','Ficha y actualización de datos · páginas 8 a 12','Activación desactivación y eliminación · páginas 13 y 14','Consulta de pagos y manejo de recibos · páginas 15 a 18','Reportes y exportaciones · páginas 19 y 20','Tarifa anual y simulador · página 21','Uso en teléfono y cierre de sesión · página 22','Solución de dificultades · página 23']:
 x=p(t);x.paragraph_format.space_after=Pt(4)

page('1 Iniciar sesión','Acceso al sistema con una cuenta de administrador.','acceso',[
'Abra la dirección de AQUAPAY indicada por su organización. Haga clic en DPI (1) e ingrese sus 13 dígitos registrados.',
'Escriba su Contraseña (2). El icono del ojo permite mostrarla u ocultarla.',
'Haga clic en Iniciar Sesión (3). Compruebe que el encabezado o menú identifique su rol como administrador.',
'En un equipo compartido, desmarque Recordarme (4). Al terminar, utilice Salir desde el menú lateral.'
],'El botón se habilita con un DPI completo y una contraseña de al menos ocho caracteres. Si sus datos son rechazados, compruebe la escritura antes de intentarlo otra vez.')
page('2 Consultar el Dashboard','Resumen de la comunidad y accesos del panel administrativo.','01_dashboard',[
'Revise los indicadores Beneficiarios, Viviendas activas, Chorros activos y Total cobrado. Este último muestra la recaudación del mes.',
'Use Actualizar resumen (1) para volver a consultar las cifras. La gráfica presenta la evolución mensual y el acumulado del año.',
'Revise Estado del padrón: activos, inactivos, personas sin vivienda y proporción de beneficiarios activos con conexión.',
'Abra Beneficiarios (2) para gestionar el padrón o Reportes (3) para consultar listados. Más abajo encontrará los últimos registros y la actividad de pagos.'
],'Si no hay cobros registrados, el resumen financiero puede mostrar cero. Eso no equivale a un error de conexión.')
page('3 Buscar y exportar beneficiarios','Controles para consultar y organizar el padrón de beneficiarios.','02_padron',[
'Escriba un nombre o DPI en Buscar beneficiarios (1) y pulse Buscar (2). Para recuperar la lista general, borre el texto y vuelva a pulsar Buscar.',
'Seleccione Todos, Activos (3) o Inactivos. Use Ordenar beneficiarios (4) para ordenar de A a Z, de Z a A o por registros recientes.',
'Pulse el nombre de una persona para abrir su ficha. Al final del listado, Anterior y Siguiente recorren páginas de hasta 10 beneficiarios.',
'Use Exportar CSV (5) para descargar todos los resultados filtrados, incluidos los de otras páginas. Para iniciar un alta, pulse Registrar (6).'
],'Los indicadores superiores describen la consulta cargada. El filtro de estado modifica el listado y la exportación.',True)
page('4 Registrar los datos de la persona','Primer paso del registro de un beneficiario.','03_persona',[
'Desde Beneficiarios, haga clic en Registrar. Complete Nombre completo (1), con un mínimo de tres caracteres.',
'Ingrese el DPI de 13 dígitos (2) y verifique que corresponda a la persona. Complete Teléfono (3) si dispone de él.',
'Revise los datos y haga clic en Continuar (4). El sistema guarda la persona y abre el paso Vivienda.',
'Si aparece un mensaje de error, corríjalo antes de seguir. Si el DPI ya está registrado, busque la ficha existente.'
],'El asistente guarda por etapas. Salir después de Continuar no elimina los datos ya creados; revise el padrón antes de iniciar otro registro.',True)
page('5 Registrar la primera vivienda','Segundo paso del registro de un beneficiario.','04_vivienda',[
'Escriba la Dirección (1) con referencias suficientes para identificar la vivienda.',
'Revise Año de inicio de cobro (2). Indique el año desde el que corresponde considerar el servicio; el formulario admite valores de 2000 a 2100.',
'Compruebe Vivienda activa (3). Mantenga el estado que corresponda a la situación del servicio.',
'Pulse Continuar (4) para guardar la vivienda y pasar a Chorro. Anterior permite revisar los datos de la persona.'
],'El año de inicio influye en la consulta de obligaciones del servicio. Verifíquelo antes de guardar.',True)
page('6 Registrar la instalación del chorro','Tercer paso del registro de un beneficiario.','05_chorro',[
'Ingrese la Cantidad de chorros (1), como número entero de al menos una unidad.',
'Revise Precio unitario (2) y Total de compra (3). El total se calcula automáticamente según la cantidad.',
'Compruebe la Fecha de instalación (4). Use Anterior si necesita revisar la vivienda.',
'Pulse Finalizar (5). Espere la confirmación y la apertura de la ficha del beneficiario; revise que aparezcan la vivienda y los chorros.'
],'Registrar el chorro no confirma que su compra esté pagada. Tesorería registra el cobro y emite el recibo correspondiente. El precio mostrado en esta versión es Q 1,500.00 por unidad.',True)
page('7 Consultar la ficha del beneficiario','Información y acciones disponibles en la ficha del beneficiario.','06_ficha',[
'En Beneficiarios, haga clic en el nombre. Compruebe identidad, DPI, estado, teléfono y fecha de registro.',
'Use Editar datos (1) para corregir nombre o teléfono. Agregar vivienda (2) permite asociar otra propiedad a esta misma persona.',
'En Viviendas y chorros, identifique la dirección correcta antes de usar Editar dirección (3) o Agregar chorro (4).',
'Use Consultar pago (5) para abrir la consulta con el DPI de esta persona. El periodo inicial es el año actual.'
],'Los recuentos de la ficha muestran registros asociados. Revise también los estados de la persona, de la vivienda y de sus chorros.',True)
page('8 Editar nombre y teléfono','Ventana para actualizar los datos de contacto del beneficiario.','07_editar',[
'En la ficha, haga clic en Editar datos.',
'Corrija Nombre completo (1) y Teléfono (2). Revise la escritura antes de confirmar.',
'Pulse Guardar cambios (3) y espere el mensaje de actualización.',
'Compruebe el resultado en la ficha. Para salir sin guardar, cierre la ventana con la X.'
],'El DPI no se puede modificar en esta ventana. Si detecta un error de identificación, revise el caso antes de eliminar o duplicar registros.')
page('9 Agregar otra vivienda','Ventana para asociar una vivienda adicional al beneficiario.','08_nueva_vivienda',[
'En la ficha de la persona correcta, haga clic en Agregar vivienda.',
'Complete Dirección (1) y Año inicio de cobro (2). Compruebe que no sea una vivienda ya registrada.',
'Pulse Guardar (3). Espere la confirmación y revise la nueva vivienda en la ficha.',
'Si corresponde instalar conexiones, use Agregar chorro en la vivienda recién creada.'
],'Para asociar otra vivienda a la misma persona no necesita registrar de nuevo su DPI.')
page('10 Corregir la dirección de una vivienda','Ventana para modificar la dirección de la vivienda seleccionada.','09_direccion',[
'Ubique la vivienda correcta dentro de la ficha y pulse Editar dirección.',
'Corrija Dirección (1). Ingrese entre 3 y 300 caracteres, con referencias claras.',
'Pulse Guardar dirección (2) y compruebe el texto actualizado.',
'Para conservar la dirección anterior, pulse Cancelar (3). Si el guardado falla, revise el mensaje de la ventana.'
],'Esta ventana modifica la dirección. Revise que la corrección corresponda a la misma vivienda.')
page('11 Agregar chorros a una vivienda','Ventana para registrar nuevas conexiones en una vivienda existente.','10_nuevo_chorro',[
'En la ficha, localice la vivienda y haga clic en Agregar chorro.',
'Indique Cantidad de chorros (1). Revise el precio unitario y el Precio total de compra (2).',
'Pulse Guardar (3). Espere la confirmación y compruebe que la instalación aparezca bajo esa vivienda.',
'Coordine con Tesorería el registro del pago de la compra cuando corresponda.'
],'El importe de compra se calcula automáticamente. Esta operación registra la conexión y no emite un recibo de pago.')
page('12 Desactivar y reactivar beneficiarios','Confirmación para cambiar la disponibilidad de un beneficiario.','11_desactivar',[
'Desde el padrón o la ficha, pulse el interruptor del estado de la persona.',
'Revise el nombre y DPI en la confirmación. Si el cambio no corresponde, pulse Cancelar (1).',
'Para desactivar, pulse Sí, desactivar (2) y compruebe el estado Inactivo.',
'Para reactivar, busque a la persona entre Inactivos, pulse su interruptor y confirme Sí, activar.'
],'Un beneficiario inactivo no se puede seleccionar para nuevos cobros. La desactivación conserva sus datos e historial y permite volver a activarlo.')
page('13 Eliminar un beneficiario','Confirmación de eliminación del beneficiario y sus registros asociados.','12_eliminar_beneficiario',[
'En la ficha pulse Eliminar beneficiario, o use Eliminar en la fila correspondiente del padrón.',
'Revise la identidad y el alcance de la eliminación. Para conservar los registros, pulse Cancelar (2).',
'Si procede eliminar, marque la casilla También eliminar sus recibos (1). Esta aceptación habilita Sí, eliminar (3).',
'Pulse Sí, eliminar y espere la confirmación. Revise que el beneficiario ya no aparezca en el padrón.'
],'La eliminación no se puede deshacer: borra al beneficiario, viviendas, chorros, recibos y pagos asociados. El siguiente correlativo se calcula como el mayor recibo restante del año más uno, o 001 si no queda ninguno. Los recibos conservados mantienen sus números.')
page('14 Buscar pagos registrados','Búsqueda de pagos agrupados por beneficiario.','13_pagos',[
'Abra Consultar pago en el menú. En la búsqueda (1), escriba el nombre o DPI y pulse Buscar (3).',
'Seleccione Año (2). Cambiar el año ejecuta la consulta; Todos los años amplía el periodo. Limpiar devuelve la búsqueda al año actual y borra el texto.',
'Revise Personas con pago, Recibos en la consulta e Importe consultado. El listado agrupa recibos por persona.',
'Pulse Ver recibos (4) junto a la persona para desplegar sus comprobantes.'
],'No aparecer en esta consulta solo indica que no se encontraron pagos con esos filtros. No equivale, por sí solo, a una certificación de deuda. La consulta solicita hasta 500 recibos.',True)
page('15 Revisar los recibos de una persona','Detalle de los recibos asociados a la persona seleccionada.','14_recibos',[
'Después de Ver recibos, revise número, fecha, tipo de pago, total y tesorero.',
'Pulse Ver PDF (1) para consultar el comprobante, descargarlo o imprimirlo.',
'Use Abrir ficha (3) para consultar los datos de la persona. Ocultar contrae la sección de recibos.',
'Eliminar (2) abre una confirmación para el recibo de esa misma fila. Revise cuidadosamente el número antes de continuar.'
],'La consulta administrativa permite revisar comprobantes existentes. La emisión de nuevos cobros se realiza desde Tesorería.',True)
page('16 Descargar o imprimir un recibo','Controles del visor de un recibo con documento de demostración.','15_pdf',[
'Pulse Ver PDF en el recibo que necesita. Espere a que cargue la vista previa y compruebe su número.',
'Pulse Imprimir (1). Seleccione impresora y opciones en el diálogo del navegador y confirme la impresión.',
'Pulse Descargar (2) para guardar el PDF. Revise la carpeta de descargas o la ubicación seleccionada.',
'Pulse Cerrar (3) para regresar a la consulta de pagos.'
],'Abrir, descargar o imprimir un recibo existente no registra un nuevo pago. Si el visor falla, cierre la ventana y vuelva a abrir el mismo comprobante.')
page('17 Eliminar un recibo','Confirmación de eliminación de un comprobante específico.','16_eliminar_recibo',[
'En Consultar pago, despliegue Ver recibos de la persona y pulse Eliminar en el comprobante correcto.',
'Revise número, beneficiario e importe en la ventana de confirmación.',
'Si no corresponde eliminarlo, pulse Cancelar (1).',
'Para confirmar, pulse Sí, eliminar recibo (2). Espere el mensaje de éxito y revise el listado actualizado.'
],'Esta acción elimina el recibo, sus pagos asociados y su PDF. El cuadro no ofrece una opción de deshacer. Verifique el registro antes de confirmar.')
page('18 Consultar y exportar reportes','Categorías y herramientas del Centro de reportes.','17_reportes',[
'Abra Reportes y seleccione Usuarios (1), Viviendas (2), Recibos (3) o Ingresos (4). Usuarios es la categoría inicial.',
'Escriba en Buscar dentro de este reporte (5) para filtrar sus filas. Borre el texto para volver a mostrar todas.',
'Use las flechas al final de la tabla para recorrer páginas de hasta 12 registros. Los indicadores superiores describen el reporte completo cargado.',
'Pulse Exportar CSV (6) para descargar los resultados filtrados, o Exportar PDF (7) para descargar el reporte completo de la categoría y periodo.'
],'La búsqueda de texto afecta el CSV y la tabla; no recorta el PDF. Usuarios y Viviendas incluyen todos los registros del padrón.',True)
page('19 Aplicar el año de un reporte','Selección del periodo para reportes de Recibos e Ingresos.','18_anual',[
'Seleccione Recibos o Ingresos. Estas categorías muestran un periodo anual.',
'Escriba el año en Año del reporte (1), dentro del intervalo de 2000 a 2100.',
'Pulse Aplicar año (2). Compruebe el periodo visible y espere a que se actualicen los datos.',
'Descargue el reporte con Exportar PDF (3) o Exportar CSV (4). Si utiliza búsqueda de texto, recuerde que solo limita el CSV.'
],'Escribir otro año no lo aplica hasta pulsar Aplicar año. Un periodo sin movimientos puede mostrar el estado Sin movimientos en este período.',True)
page('20 Actualizar la tarifa anual','Tarifa vigente y simulación de una propuesta antes de guardar.','19_tarifa',[
'Abra Tarifa. Consulte el importe de Tarifa vigente y cambie Tarifa anual por chorro (1) solo cuando corresponda actualizarlo.',
'En Cantidad de chorros (2), pruebe el cálculo del Simulador de tarifa. La estimación usa el importe escrito, aunque todavía no esté guardado.',
'Revise la propuesta y pulse Guardar cambios (3). Se habilita al introducir un importe distinto del vigente y mayor o igual que Q 0.01.',
'Espere el mensaje Tarifa anual actualizada y compruebe el nuevo importe vigente. Lugar de pago (4) es un campo de consulta.'
],'El simulador no registra cobros. La tarifa vigente se conserva hasta guardar. La compra de chorro es un concepto separado y no está incluida en la tarifa anual.',True)
page('21 Usar el teléfono y cerrar la sesión','Menú administrativo en una pantalla de teléfono.','20_movil',[
'Pulse Abrir menú en el encabezado. Toque Beneficiarios (1), Reportes (2) u otro módulo para entrar.',
'En las tablas anchas, deslice horizontalmente para ver las columnas y acciones. Desplácese hacia abajo para encontrar los controles de página.',
'Al terminar, abra el menú y pulse Salir (3). Compruebe el regreso a la pantalla de acceso.',
'El sistema también cierra la sesión después de cinco minutos sin actividad. Si ocurre durante un registro, vuelva a entrar y compruebe qué pasos quedaron guardados antes de repetirlos.'
],'En un equipo compartido, cierre la sesión aunque haya desmarcado Recordarme. Proteja los archivos descargados según las reglas de su organización.')
page('22 Resolver dificultades','Acciones recomendadas ante problemas frecuentes del panel.')
for title,text in [
('No aparece un beneficiario','Borre la búsqueda, pulse Buscar y seleccione Todos. Compruebe el DPI, el nombre y las páginas del listado. Un registro inactivo puede quedar oculto por el filtro Activos.'),
('El registro se interrumpió antes de Finalizar','Busque el DPI en el padrón. La persona puede haberse guardado en el primer paso y la vivienda en el segundo. Abra la ficha y complete la vivienda o los chorros que falten, evitando duplicados.'),
('El sistema rechaza el DPI o un formulario','Compruebe los 13 dígitos del DPI, los campos obligatorios y el mensaje mostrado. En el registro, el DPI queda bloqueado después de guardar a la persona; en Editar datos no se modifica.'),
('Guardar cambios de tarifa está desactivado','Espere a que termine la carga. Introduzca un valor numérico de al menos Q 0.01 y distinto del vigente. El simulador no guarda automáticamente.'),
('El reporte muestra el año anterior a mi cambio','Pulse Aplicar año y compruebe el periodo del encabezado. Al exportar PDF se utiliza el año aplicado, no un valor escrito que aún no se haya confirmado.'),
('No hay pagos o el PDF no carga','Revise nombre, DPI y año de pago. Si no hay resultados, amplíe los filtros. Si el recibo existe pero su PDF no abre, cierre el visor y vuelva a intentarlo sin registrar un pago nuevo.'),
('Aparece un error de conexión o de guardado','Verifique internet y use Reintentar cuando esté disponible. Antes de repetir un alta o una eliminación, consulte la ficha o el listado para comprobar si la operación llegó a completarse.')]:
 doc.add_heading(title,2);p(text)
p('Al solicitar soporte, indique el módulo, la acción, el mensaje exacto y el registro afectado. No comparta su contraseña.')
doc.save(OUT);print(OUT)
