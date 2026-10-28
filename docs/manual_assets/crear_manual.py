from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parent/'Manual_de_Usuario_Tesoreria_AQUAPAY.docx'
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
for s in doc.styles:
 for b in list(s.element.findall('.//' + qn('w:pBdr'))):
  b.getparent().remove(b)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(8)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
doc.styles['Heading 1'].font.size=Pt(20)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(12)
doc.styles['Title'].font.size=Pt(30)
doc.styles['Caption'].font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('AQUAPAY  |  Tesorería  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='Manual de usuario del panel de Tesorería de AQUAPAY'
doc.core_properties.author='AQUAPAY'

def p(text):return doc.add_paragraph(text)
def steps(items):
 for n,t in enumerate(items,1):
  x=doc.add_paragraph();x.add_run(f'{n}. ').bold=True;x.add_run(t)
def img(name,caption,crop=False,width=7.1):
 f=ROOT/(name+'.png')
 if name=='01_acceso':
  im=Image.open(f).crop((680,210,1250,710));f=ROOT/'01_acceso_detalle.png';im.save(f);width=5.2
 if name=='09_eliminar':
  im=Image.open(f).crop((435,265,925,625));f=ROOT/'09_eliminar_detalle.png';im.save(f);width=5.2
 if crop:
  im=Image.open(f);im=im.crop((250,75,1360,900));f=ROOT/(name+'_detalle.png');im.save(f)
 x=doc.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_after=Pt(4)
 x.add_run().add_picture(str(f),width=Inches(width))
 # Describe the figure for assistive readers.
 doc.inline_shapes[-1]._inline.docPr.set('descr',caption)
 doc.add_paragraph(caption,'Caption')
def page(title,lead,name=None,caption='',items=[],note='',crop=False):
 doc.add_page_break();doc.add_heading(title,1);p(lead)
 if name:img(name,caption,crop)
 steps(items)
 if note:
  x=p('');x.add_run('Tenga en cuenta ').bold=True;x.add_run(note)

doc.add_paragraph('AQUAPAY','Subtitle')
doc.add_paragraph('Manual de usuario del panel de Tesorería','Title')
p('Guía de operación para el personal responsable de registrar cobros del servicio de agua y entregar recibos a los beneficiarios.')
p('Versión 1.0  |  27 de octubre de 2026')
doc.add_heading('Cómo utilizar este manual',1)
p('Busque la tarea que necesita realizar y siga los pasos en orden. Los círculos rojos sobre las capturas señalan los campos y botones mencionados en las instrucciones. Los nombres, DPI, importes y recibos de ejemplo sirven para practicar la lectura de las pantallas; en su sesión aparecerán los datos de su comunidad.')
p('Antes de empezar, disponga de la dirección de AQUAPAY proporcionada por el administrador, conexión a internet y una cuenta activa con rol de tesorero. Para imprimir, tenga una impresora disponible o seleccione la opción de guardar como PDF de su navegador.')
doc.add_heading('Contenido',1)
for text in ['1  Iniciar sesión · página 2','2  Consultar el resumen · página 3','3  Buscar un beneficiario · página 4','4  Verificar la vivienda · página 5','5  Cobrar la tarifa anual · página 6','6  Cobrar una compra de chorro · página 7','7  Consultar y exportar el historial · página 8','8  Descargar o imprimir un recibo · página 9','9  Eliminar un recibo · página 10','10  Recorrer los resultados · página 11','11  Cerrar la sesión · página 12','12  Usar el panel en el teléfono · página 13','13  Resolver dificultades · página 14']:
 x=p(text);x.paragraph_format.space_after=Pt(3)

page('1 Iniciar sesión','Abra la dirección de AQUAPAY en su navegador. Use los datos de su cuenta de Tesorería.','01_acceso','Figura 1. Acceso al sistema y controles de identificación.',[
'Haga clic en DPI (1) e ingrese los 13 dígitos de su documento registrado.',
'Haga clic en Contraseña (2) y escriba su clave. El icono del ojo permite mostrarla u ocultarla.',
'Haga clic en Iniciar Sesión (3). Al validar los datos se abrirá el Dashboard.',
'Revise Recordarme (4). En un equipo compartido, desmarque esta casilla y cierre siempre la sesión al terminar.'
], 'El botón de acceso requiere un DPI completo y una contraseña de al menos ocho caracteres. Si sus credenciales son rechazadas, verifique los datos o solicite ayuda al administrador.')
page('2 Consultar el resumen','El Dashboard presenta el resumen de Tesorería del mes y la evolución de ingresos del año.','02_inicio','Figura 2. Resumen de Tesorería y accesos de trabajo.',[
'Consulte Ingresos del mes, Tarifa anual, Compra de chorro y Recibos. Las dos tarjetas por concepto muestran recaudación, mientras que Servicio de agua muestra la tarifa anual vigente por chorro.',
'Para registrar un cobro, haga clic en Emitir (1). Para refrescar el resumen, use el botón con flecha circular (2).',
'Revise Evolución de ingresos y el acumulado del año. En Servicio de agua consulte beneficiarios y viviendas activos.',
'Para consultar comprobantes anteriores, haga clic en Historial (3) o en Ver historial dentro de Últimos recibos. Ver PDF abre un recibo sin emitir otro.'
])
page('3 Buscar un beneficiario','Entre en Emitir desde el menú lateral. El primer paso del formulario identifica a quién corresponde el cobro.','03_buscar','Figura 3. Búsqueda y selección del beneficiario.',[
'Haga clic en el campo de búsqueda (1). Escriba parte del nombre o el DPI del beneficiario.',
'Espere la lista de coincidencias. Se muestran beneficiarios activos.',
'Haga clic en el resultado correcto (2). Escribir el nombre sin seleccionar el resultado no completa la elección.',
'Compruebe los datos cargados antes de continuar. Si cambia la búsqueda después de seleccionar, deberá elegir nuevamente al beneficiario.'
], 'Si no encuentra al beneficiario, revise la escritura o pruebe con el DPI. Si sigue sin aparecer, solicite al administrador que compruebe el registro y su estado activo.',True)
page('4 Verificar la vivienda','La elección del beneficiario carga sus datos y consulta los pagos pendientes de la vivienda activa.','04_vivienda','Figura 4. Datos del beneficiario y de la vivienda seleccionada.',[
'Revise Nombre (1) y DPI (2). Estos campos se completan automáticamente.',
'Verifique Dirección (3). Si el beneficiario tiene varias viviendas, seleccione la correcta en Dirección / vivienda; la consulta de pendientes cambiará con la selección.',
'Compruebe Cantidad de chorros (4). Este dato participa en el cálculo de la tarifa anual.',
'Espere a que termine Consultando pendientes. Si aparece un error, haga clic en Reintentar consulta antes de continuar.'
], 'El tesorero no modifica estos datos desde el formulario de cobro. Si falta una vivienda activa o la dirección es incorrecta, solicite la corrección al administrador.',True)
page('5 Cobrar la tarifa anual','Seleccione el concepto y revise el resumen antes de registrar el dinero recibido.','05_tarifa','Figura 5. Tarifa anual, periodo a cancelar y resumen del cobro.',[
'Seleccione Tarifa anual (1). En Año a cancelar (2), escriba el año hasta el que pagará el beneficiario.',
'Revise los Años incluidos en el resumen. El sistema suma los años pendientes hasta el año indicado. Por ejemplo, si debe 2024, 2025 y 2026 y selecciona 2026, se incluyen los tres años.',
'Revise la Fecha de pago (3) y, si lo necesita, agregue una Descripción (4). Se permiten hasta 500 caracteres y se usa una descripción automática si deja el campo vacío.',
'Verifique beneficiario, vivienda, concepto, años, fecha y Total a pagar. Solo después de recibir el pago, haga clic en Emitir (5). Espere a que se abra el visor PDF.'
], 'El año permitido va de 2000 al año actual. El correlativo mostrado es preliminar: el número definitivo se asigna al guardar. No vuelva a pulsar Emitir mientras aparezca Generando.',True)
page('6 Cobrar una compra de chorro','Esta opción registra una compra pendiente asociada a la vivienda seleccionada.','06_compra','Figura 6. Selección de una compra de chorro pendiente.',[
'Seleccione Compra de chorro (1).',
'En Chorro a cobrar (2), elija la compra pendiente correspondiente y compruebe la cantidad de unidades y el importe.',
'Revise Fecha de pago (3) y complete la descripción si necesita una observación en el recibo.',
'Compare los datos del Resumen del cobro con el dinero recibido. Haga clic en Emitir (4) y espere el visor del recibo.'
], 'Si aparece Sin compra pendiente, no hay una compra disponible para esa vivienda. Solicite al administrador que revise el registro si considera que falta una. Tras emitir correctamente, el formulario se limpia y vuelve a Tarifa anual.',True)
page('7 Consultar y exportar el historial','Entre en Historial para buscar comprobantes y revisar los totales de la consulta.','07_historial','Figura 7. Filtros de búsqueda y exportación del historial.',[
'En Beneficiario o recibo (1), escriba nombre, DPI o número de recibo. Puede dejarlo vacío para ampliar la consulta.',
'Seleccione Año de pago (2) y Concepto (3), según lo que necesita. El año corresponde a la fecha del pago, no necesariamente al año de tarifa cancelado.',
'Haga clic en Buscar (4) para aplicar el texto y el año. Concepto filtra la lista cargada. Para volver a la consulta general, use Limpiar filtros.',
'Revise los indicadores y haga clic en Exportar CSV (5). El navegador descargará un archivo que puede abrir en Excel; contiene todos los resultados filtrados de la consulta, no solo la página visible.'
], 'La consulta admite como máximo 500 recibos. Si aparece el aviso del límite, acote por año o beneficiario. Los totales y el CSV corresponden a los resultados cargados y filtrados.',True)
page('8 Descargar o imprimir un recibo','En Últimos recibos o en Historial, localice el comprobante y haga clic en Ver PDF.','08_pdf','Figura 8. Controles del visor PDF con un documento de demostración.',[
'Espere a que cargue la vista previa. Verifique que el número y el contenido correspondan al pago que necesita.',
'Haga clic en Imprimir (1). En el diálogo del navegador, seleccione la impresora, revise las opciones y confirme la impresión.',
'Haga clic en Descargar (2) para guardar el archivo PDF. Revise la carpeta de descargas o la ubicación elegida en su navegador.',
'Haga clic en Cerrar (3) para regresar al panel. Abrir, imprimir o descargar un recibo existente no registra un nuevo cobro.'
], 'Si el PDF no carga después de emitir, consulte primero el Historial y abra el mismo recibo. Evite registrar nuevamente un pago para recuperar su comprobante.')
page('9 Eliminar un recibo','Use esta función únicamente cuando corresponda corregir un registro y tenga autorización de su organización.','09_eliminar','Figura 9. Confirmación de eliminación del recibo.',[
'En Historial o Últimos recibos de Emitir, ubique el comprobante y haga clic en Eliminar en su misma fila.',
'En la ventana Eliminar recibo, revise el número, el beneficiario y el importe.',
'Si no desea continuar, haga clic en Cancelar (1). El recibo se conserva.',
'Si verificó que procede la eliminación, haga clic en Sí, eliminar recibo (2). Espere el mensaje de confirmación y revise el listado actualizado.'
], 'La acción elimina el recibo, sus pagos asociados y su PDF. El cuadro no ofrece una opción de deshacer; confirme cuidadosamente el registro antes de continuar.')
page('10 Recorrer los resultados','El historial organiza los resultados en páginas de hasta 12 recibos.','10_paginas','Figura 10. Controles de paginación al final del listado.',[
'Desplácese al final del listado y revise Mostrando para conocer el intervalo visible y el total de resultados.',
'Haga clic en Siguiente (2) para avanzar a los recibos de la siguiente página.',
'Haga clic en Anterior (1) para regresar. Un botón desactivado indica que llegó al inicio o al final.',
'Para recuperar otro periodo, ajuste los filtros y pulse Buscar. La consulta volverá a la primera página.'
], 'Cambiar de página no modifica los recibos. La exportación CSV y los indicadores consideran la lista filtrada completa de la consulta, aunque solo vea una parte en pantalla.',True)
page('11 Cerrar la sesión','Al terminar sus operaciones, cierre la sesión desde el menú lateral.','11_sesion','Figura 11. Botón Salir en la parte inferior del menú.',[
'Compruebe que terminó la emisión y que guardó o imprimió los comprobantes necesarios.',
'Haga clic en Salir (1), en la parte inferior del menú lateral.',
'Verifique que el sistema regrese a la pantalla de inicio de sesión.',
'En un equipo compartido, cierre el navegador después de salir y proteja los archivos descargados según las reglas de su organización.'
], 'El sistema cierra la sesión después de cinco minutos sin actividad. Si ocurre, vuelva a ingresar y compruebe en Historial si el último recibo se registró antes de intentar emitir de nuevo.')
doc.add_page_break();doc.add_heading('12 Usar el panel en el teléfono',1)
p('En una pantalla pequeña, el menú se abre con el botón Abrir menú del encabezado. Las funciones mantienen los mismos nombres.')
img('12_movil','Figura 12. Menú de Tesorería en un teléfono.',width=2.55)
steps(['Abra el menú y toque Emitir (1) o Historial (2).','En Emitir, complete beneficiario, vivienda, concepto y fecha; desplácese hacia abajo para revisar el resumen y tocar Emitir.','En Historial, los recibos se presentan como tarjetas con Ver PDF y Eliminar. Para terminar, abra el menú y toque Salir (3).'])
page('13 Resolver dificultades','Identifique el mensaje en pantalla y realice la acción correspondiente.')
for title,text in [
('El botón Emitir está desactivado','Seleccione un beneficiario de la lista, verifique la vivienda activa y espere la consulta de pendientes. Revise que haya una fecha y un total mayor que cero. En tarifa anual, corrija el año o compruebe que exista deuda; en compra de chorro, elija una compra pendiente.'),
('No hay tarifa anual pendiente','La vivienda no tiene deuda anual disponible en la consulta. No intente forzar el cobro cambiando el importe. Si el dato parece incorrecto, solicite una revisión al administrador.'),
('No se pudieron cargar los datos','Revise la conexión a internet y use Reintentar o Reintentar consulta. Un mensaje de error no significa que no existan beneficiarios o pagos.'),
('No hay coincidencias en el historial','Revise nombre, DPI, recibo, año de pago y concepto. Use Limpiar filtros o Ver todos los recibos para ampliar la búsqueda.'),
('Hubo un error o una interrupción al emitir','El formulario conserva los datos cuando falla el cobro. Antes de repetirlo, revise Historial para comprobar si el recibo llegó a registrarse. Si ya existe, abra su PDF. Si no puede confirmar el resultado, solicite ayuda al administrador.'),
('No se muestra el PDF','Cierre el visor, vuelva a buscar el recibo en Historial y pulse Ver PDF. Si persiste el problema, informe al administrador el número del recibo y el mensaje mostrado.'),
('El total o los datos del beneficiario no corresponden','Detenga la emisión y compruebe vivienda, cantidad de chorros, años incluidos y tipo de pago. La tarifa y los registros de beneficiarios se revisan con el administrador.')]:
 doc.add_heading(title,2);p(text)
p('Para pedir ayuda, indique la tarea que realizaba, el número de recibo si existe y el mensaje exacto. No comparta su contraseña.')
doc.save(OUT);print(OUT)

